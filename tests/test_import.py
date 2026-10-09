"""Імпорт початкових залишків з Excel (шаблон програми)."""
from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from oblik import balances, documents, importer, locations
from oblik.documents import DocumentError
from tests.test_documents_web import add, env, open_db  # noqa: F401 — фікстура env

HEAD = [c[0] for c in importer.COLUMNS]


def xlsx(rows, head=HEAD, sheet=importer.SHEET) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(head)
    for r in rows:
        ws.append(r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


GOOD = [
    # наявне найменування (Кабель, м), поштучне наявне (Камера), нове кількісне і нове поштучне
    ["Кабель", "", "м", "", "", "", "", "I", "10,50", 250, ""],
    ["Камера", "", "шт.", "", "", "SN-001", "", "II", 4460, 1, ""],
    ["Батарейка АА", "1234", "шт.", "Запаси", "", "", "", "", 25, 40, "для пультів"],
    [None] * 11,                                                       # порожній — пропускається
    ["Реєстратор", "", "шт.", "Основні засоби", "так", "", "ІНВ-7", "I", "12 000", 1, ""],
]


def test_template_has_sheets_and_columns():
    wb = load_workbook(BytesIO(importer.template_xlsx()))
    assert wb.sheetnames == ["Залишки", "Приклад", "Як заповнювати"]
    assert [c.value.rstrip(" *") for c in wb["Залишки"][1]] == HEAD
    # Приклад — сам по собі коректний файл імпорту.
    ex = wb["Приклад"]
    assert ex.max_row == 1 + len(importer.EXAMPLE)


def test_import_creates_draft_and_refs(env):
    storage, _, ids = env
    conn = open_db(storage)
    wh = locations.for_unit(conn, ids["wh"])
    res = importer.parse(conn, xlsx(GOOD))
    assert res.errors == [] and len(res.lines) == 4
    # Приклад із шаблону — коректний файл.
    assert importer.parse(conn, xlsx(importer.EXAMPLE)).errors == []
    assert res.new_names == ["Батарейка АА", "Реєстратор"]
    doc, res = importer.create_opening(conn, xlsx(GOOD), wh, "2026-10-01")
    d = documents.get(conn, doc)
    assert (d["doc_type"], d["status"], d["op_date"], d["to_location_id"]) == ("opening", "draft", "2026-10-01", wh)
    documents.post(conn, doc)
    names = {r["id"]: r for r in conn.execute("SELECT * FROM nomenclature")}
    left = {(names[r["nomenclature_id"]]["name"], r["category"]): (r["qty_m"], r["price_kop"])
            for r in balances.balances(conn, [wh], "2026-10-31")}
    assert left == {("Кабель", 1): (250000, 1050), ("Камера", 2): (1000, 446000),
                    ("Батарейка АА", 1): (40000, 2500), ("Реєстратор", 1): (1000, 1200000)}
    bat = conn.execute("SELECT * FROM nomenclature WHERE name = 'Батарейка АА'").fetchone()
    assert bat["code"] == "1234" and bat["accounting_class"] == "inventory" and not bat["serial_tracked"]
    reg = conn.execute("SELECT * FROM nomenclature WHERE name = 'Реєстратор'").fetchone()
    assert reg["serial_tracked"] and reg["accounting_class"] == "fixed"
    assert conn.execute("SELECT COUNT(*) FROM items WHERE inventory_no = 'ІНВ-7'").fetchone()[0] == 1


def test_import_errors_change_nothing(env):
    storage, _, ids = env
    conn = open_db(storage)
    wh = locations.for_unit(conn, ids["wh"])
    noms = conn.execute("SELECT COUNT(*) FROM nomenclature").fetchone()[0]
    bad = [
        ["", "", "шт.", "", "", "", "", "", 1, 1, ""],                          # 2: без назви
        ["Кабель", "", "кг", "", "", "", "", "", 1, 1, ""],                     # 3: інша од. виміру
        ["Камера", "", "шт.", "", "", "", "", "", 1, 1, ""],                    # 4: поштучне без номера
        ["Камера", "", "шт.", "", "", "SN-9", "", "", 1, 2, ""],                # 5: поштучне, кількість 2
        ["Нове", "", "шт.", "", "", "", "", "VI", 1, 1, ""],                    # 6: категорія
        ["Нове2", "", "шт.", "Щось", "", "", "", "", 1, 1, ""],                 # 7: клас
        ["Нове3", "", "шт.", "", "", "", "", "", "", 1, ""],                    # 8: без ціни
        ["Нове4", "", "шт.", "", "", "", "", "", 1, "абв", ""],                 # 9: кількість
        ["Камера", "", "шт.", "", "", "SN-5", "", "", 1, 1, ""],                # 10
        ["Камера", "", "шт.", "", "", "SN-5", "", "", 1, 1, ""],                # 11: номер повторюється
        ["Кабель", "", "м", "", "", "SN-1", "", "", 1, 1, ""],                  # 12: номер для кількісного
    ]
    with pytest.raises(DocumentError) as exc:
        importer.create_opening(conn, xlsx(bad), wh, "2026-10-01")
    errors = exc.value.errors
    for n, text in [(2, "не вказано найменування"), (3, "обліковується в «м»"), (4, "вкажіть заводський"),
                    (5, "кількість 1"), (6, "невідома категорія"), (7, "невідомий клас"), (8, "ціну"),
                    (9, "Рядок 9:"), (11, "уже є в рядку 10"), (12, "обліковується кількістю")]:
        assert any(e.startswith(f"Рядок {n}:") and text in e for e in errors), (n, errors)
    assert conn.execute("SELECT COUNT(*) FROM nomenclature").fetchone()[0] == noms
    assert conn.execute("SELECT COUNT(*) FROM documents").fetchone()[0] == 0
    # Чужий файл / немає обов'язкової графи / порожній.
    assert "Не вдалося прочитати файл" in importer.parse(conn, b"not excel").errors[0]
    assert "немає графи «Кількість»" in importer.parse(conn, xlsx([], head=HEAD[:9])).errors[0]
    assert "немає жодного заповненого рядка" in importer.parse(conn, xlsx([])).errors[0]


def test_web_import(env):
    storage, client, ids = env
    resp = client.get("/documents/import/template.xlsx")
    assert resp.status_code == 200 and resp.data[:2] == b"PK"
    assert "/documents/import" in client.get("/documents/new/opening").get_data(as_text=True)
    page = client.get("/documents/import").get_data(as_text=True)
    assert "Завантажити шаблон" in page and "Що це за сторінка" in page
    resp = client.post("/documents/import", data={
        "to_place": f"u:{ids['wh']}", "op_date": "2026-10-01",
        "file": (BytesIO(xlsx([["Кабель", "", "кг", "", "", "", "", "", 1, 1, ""]])), "z.xlsx")},
        content_type="multipart/form-data")
    html = resp.get_data(as_text=True)
    assert resp.status_code == 200 and "Рядок 2:" in html and "нічого не змінено" in html
    resp = client.post("/documents/import", data={
        "to_place": f"u:{ids['wh']}", "op_date": "2026-10-01", "file": (BytesIO(xlsx(GOOD)), "z.xlsx")},
        content_type="multipart/form-data")
    assert resp.status_code == 302
    edit = client.get(resp.headers["Location"]).get_data(as_text=True)
    assert "рядків — 4, нових найменувань — 2" in edit and "Реєстратор" in edit

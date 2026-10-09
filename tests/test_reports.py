"""Звіти: залишки по МВО / підрозділу / особі на дату, Excel, картка руху одиниці."""
from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from oblik import db, reports
from oblik.config import Paths
from oblik.reports import ReportError
from oblik.storage import Storage
from oblik.web import create_app
from tests import test_posting
from tests.test_posting import World

# Особа рахується за ротою (місце за замовчуванням), а місць обліку два — тому з поміткою.
SOLDIER = "Боєць І. (видано: 1 рота)"


@pytest.fixture
def w(tmp_path):
    conn = db.connect(tmp_path / "t.db")
    db.migrate(conn)
    db.init_meta(conn, "test")
    world = World(conn)
    test_posting.W = world
    cam = test_posting.W.camera
    # 01.10: надійшли кабель і дві камери на склад; 05.10: камеру 1 видали бійцю.
    world.doc("receipt", world.ext, world.wh, [
        test_posting.cable(100, price=1244),
        {"nomenclature_id": cam, "item_id": world.cam1, "qty_m": 1000, "price_kop": 446000, "category": 1},
        {"nomenclature_id": cam, "item_id": world.cam2, "qty_m": 1000, "price_kop": 446000, "category": 1},
    ], op_date="2026-10-01")
    world.doc("transfer", world.wh, world.person, [
        {"nomenclature_id": cam, "item_id": world.cam1, "qty_m": 1000, "price_kop": 446000, "category": 1},
    ], op_date="2026-10-05")
    world.doc("transfer", world.wh, world.coy, [test_posting.cable(30, price=1244)], op_date="2026-10-06")
    yield world
    conn.close()


def names(report):
    return [(r["name"], r["place"], r["qty_m"]) for r in report.rows]


def test_stock_whole_unit_on_dates(w):
    rep = reports.stock(w.conn, "all", "2026-10-31")
    assert rep.many_places
    assert rep.total_kop == 100 * 1244 + 2 * 446000
    assert ("Камера", SOLDIER, 1000) in names(rep)
    # На 02.10 обидві камери ще на складі, кабель ще не передано в роту.
    early = reports.stock(w.conn, "all", "2026-10-02")
    assert {(n, p) for n, p, _ in names(early)} == {("Кабель", "Склад ТЗО"), ("Камера", "Склад ТЗО")}
    # До надходження — нічого.
    assert reports.stock(w.conn, "all", "2026-09-30").rows == []


def test_stock_by_unit_person_mvo(w):
    # Склад: кабель 70 + камера 2; камера 1 у бійця, але боєць — з роти, а не зі складу.
    wh = reports.stock(w.conn, f"u:{w.wh_unit}", "2026-10-31")
    assert sorted(names(wh)) == [("Кабель", "Склад ТЗО", 70000), ("Камера", "Склад ТЗО", 1000)]
    # Рота: кабель 30 + камера 1 у бійця (видане особі рахується в підрозділі).
    coy = reports.stock(w.conn, f"u:{w.coy_unit}", "2026-10-31")
    assert sorted(names(coy)) == [("Кабель", "1 рота", 30000), ("Камера", SOLDIER, 1000)]
    camera_row = next(r for r in coy.rows if r["item_id"])
    assert camera_row["serial_no"] == "CAM-1"
    # Особа.
    person = reports.stock(w.conn, f"p:{w.soldier}", "2026-10-31")
    assert names(person) == [("Камера", SOLDIER, 1000)] and not person.many_places
    # МВО обох місць обліку — усе майно.
    mvo = reports.stock(w.conn, f"m:{w.mvo}", "2026-10-31")
    assert mvo.total_kop == reports.stock(w.conn, "all", "2026-10-31").total_kop
    # Особа без руху — порожньо; помилки вибору — зрозумілі.
    assert reports.stock(w.conn, f"p:{w.mvo}", "2026-10-31").rows == []
    with pytest.raises(ReportError, match="не є МВО"):
        reports.stock(w.conn, f"m:{w.soldier}", "2026-10-31")
    with pytest.raises(ReportError):
        reports.stock(w.conn, f"u:{w.platoon}", "2026-10-31")  # не місце обліку


def test_stock_xlsx(w):
    rep = reports.stock(w.conn, "all", "2026-10-31")
    ws = load_workbook(BytesIO(reports.stock_xlsx(w.conn, rep))).active
    values = [[c for c in row] for row in ws.iter_rows(values_only=True)]
    assert values[1][0] == "Залишки майна на 31.10.2026"
    assert "Де перебуває" in values[4]
    flat = [v for row in values for v in row if v is not None]
    assert "CAM-1" in flat and SOLDIER in flat
    assert 1244 * 100 / 100 + 2 * 4460 in flat   # разом, грн


def test_item_card(w):
    card = reports.item_card(w.conn, w.cam1)
    assert [m["to_title"] for m in card["moves"]] == ["Склад ТЗО", SOLDIER]
    assert card["position"]["place"] == SOLDIER
    assert reports.item_card(w.conn, w.cam2)["position"]["place"] == "Склад ТЗО"
    found = reports.find_items(w.conn, "cam-1")
    assert [f["id"] for f in found] == [w.cam1] and found[0]["place"] == SOLDIER
    with pytest.raises(ReportError):
        reports.item_card(w.conn, 9999)


def test_report_pages(tmp_path):
    storage = Storage(Paths(root=tmp_path / "Флешка", local=tmp_path / "ПК"))
    storage.open()
    conn = db.connect(storage.work_db)
    world = World(conn)
    test_posting.W = world
    world.doc("receipt", world.ext, world.wh, [
        {"nomenclature_id": world.camera, "item_id": world.cam1, "qty_m": 1000, "price_kop": 100,
         "category": 2}])
    conn.commit()
    conn.close()
    app = create_app(storage)
    app.config["TESTING"] = True
    client = app.test_client()

    for url in ["/reports", "/reports/stock", f"/reports/stock?scope=u:{world.wh_unit}&on_date=2026-10-31",
                "/reports/items", "/reports/items?q=CAM", f"/reports/items/{world.cam1}"]:
        assert client.get(url).status_code == 200, url
    html = client.get("/reports/stock?scope=all&on_date=2026-10-31").get_data(as_text=True)
    assert "CAM-1" in html and "Зберегти в Excel" in html
    assert "не знайдено" in client.get("/reports/stock?scope=u:9999").get_data(as_text=True)
    resp = client.get("/reports/stock.xlsx?scope=all&on_date=2026-10-31")
    assert resp.status_code == 200 and resp.data[:2] == b"PK"
    assert client.get("/reports/items/9999").status_code == 404
    assert "Склад ТЗО" in client.get(f"/reports/items/{world.cam1}").get_data(as_text=True)


def test_stock_xlsx_has_subunit_name(w):
    w.conn.execute("INSERT INTO settings(key, value) VALUES ('unit_name', 'Військова частина А0000'), "
                   "('subunit_name', 'Військова частина А1111')")
    rep = reports.stock(w.conn, "all", "2026-10-31")
    ws = load_workbook(BytesIO(reports.stock_xlsx(w.conn, rep))).active
    first = [r[0] for r in ws.iter_rows(max_row=3, values_only=True)]
    assert first == ["Військова частина А0000", "Військова частина А1111", "Залишки майна на 31.10.2026"]

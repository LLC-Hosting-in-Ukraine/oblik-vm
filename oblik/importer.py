"""Імпорт початкових залишків з Excel (шаблон програми) → чернетка «Введення початкових залишків».

Шаблон простий і однаковий для будь-якої частини: найменування, од. виміру, клас обліку,
поштучно чи ні, номери, категорія, ціна, кількість. Файл перевіряється весь одразу: якщо є хоч
одна помилка — нічого не створюється, користувач бачить усі помилки з номерами рядків Excel.
Без помилок — нові найменування й поштучні одиниці додаються в довідники, а рядки потрапляють
у чернетку, яку можна переглянути, виправити й провести (або анулювати).

© 2026 Yevhenii Hosting by LLC Hosting in Ukraine
"""
from __future__ import annotations

import sqlite3
from dataclasses import dataclass, field
from io import BytesIO

from . import catalog, documents, locations
from .documents import DocumentError
from .textutil import parse_money, parse_qty

SHEET = "Залишки"
# (заголовок у файлі, ключ, обов'язкове, підказка в примітці до заголовка)
COLUMNS = [
    ("Найменування", "name", True, "Як у книзі обліку. Якщо таке вже є в довіднику «Номенклатура» — "
                                    "рядок прив'яжеться до нього, інакше найменування буде додано."),
    ("Код номенклатури", "code", False, "Необов'язково."),
    ("Одиниця виміру", "uom", True, "шт., м, кг, к-т, л…"),
    ("Клас обліку", "class", False, "Основні засоби / МНМА / Запаси / МШП / Забалансове. Порожньо — Запаси."),
    ("Поштучно", "serial", False, "«так» — облік за заводськими (інвентарними) номерами: один рядок — "
                                  "одна одиниця, кількість 1. Порожньо — «ні»."),
    ("Заводський №", "serial_no", False, "Для поштучного майна — заводський або інвентарний номер."),
    ("Інвентарний №", "inventory_no", False, ""),
    ("Категорія", "cat", False, "I–V (або 1–5). Порожньо — I."),
    ("Ціна за одиницю, грн", "price", True, "Наприклад: 4460 або 4460,50."),
    ("Кількість", "qty", True, "Для поштучного — 1."),
    ("Примітка", "note", False, ""),
]
CLASS_BY_TITLE = {"основні засоби": "fixed", "ос": "fixed", "мнма": "low_value",
                  "малоцінні необоротні": "low_value", "запаси": "inventory", "мшп": "mshp",
                  "забалансове": "off_balance"}
CLASS_TITLES = ["Основні засоби", "МНМА", "Запаси", "МШП", "Забалансове"]
CATEGORIES = {"i": 1, "ii": 2, "iii": 3, "iv": 4, "v": 5, "1": 1, "2": 2, "3": 3, "4": 4, "5": 5}
YES = {"так", "т", "yes", "y", "1", "+", "поштучно"}

EXAMPLE = [
    ["Камера відеоспостереження", "", "шт.", "Основні засоби", "так", "SN-001", "", "I", 4460, 1, ""],
    ["Камера відеоспостереження", "", "шт.", "Основні засоби", "так", "SN-002", "", "II", 4460, 1, ""],
    ["Кабель КПП-ВП 2х0,5", "", "м", "Запаси", "", "", "", "I", 10.5, 250, ""],
    ["Батарейка АА", "", "шт.", "Запаси", "", "", "", "", 25, 40, "для пультів"],
]

HELP = [
    "Як заповнювати шаблон",
    "",
    "1. Заповнюйте аркуш «Залишки»: один рядок — одне найменування (або одна поштучна одиниця).",
    "2. Обов'язкові графи: Найменування, Одиниця виміру, Ціна за одиницю, Кількість.",
    "3. Поштучне майно (з заводськими / інвентарними номерами): у графі «Поштучно» — «так»,",
    "   кожна одиниця — окремим рядком, кількість 1, номер — у «Заводський №» або «Інвентарний №».",
    "4. Одне найменування з різними цінами чи категоріями — різні рядки.",
    "5. Категорія — I…V (п. 7 розд. I Інструкції). Порожньо — I.",
    "6. Не змінюйте назви граф у першому рядку. Порожні рядки пропускаються.",
    "7. У програмі: «Що сталося?» → «Вносимо майно, яке вже є» → «Завантажити з Excel»: оберіть місце",
    "   обліку і дату. Програма перевірить файл; якщо є помилки — покаже їх з номерами рядків і нічого",
    "   не змінить. Без помилок — створить чернетку: перегляньте її і проведіть.",
    "",
    "Приклад заповнення — на аркуші «Приклад» (дані вигадані).",
]


def template_xlsx() -> bytes:
    """Порожній шаблон з підказками, списками вибору й прикладом."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="E8F0E0")
    widths = [38, 16, 12, 18, 11, 18, 16, 11, 14, 11, 24]
    for col, (title, _key, required, tip) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=title + (" *" if required else ""))
        cell.font, cell.fill = bold, fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        if tip:
            cell.comment = Comment(tip, "Облік ВМ")
        ws.column_dimensions[cell.column_letter].width = widths[col - 1]
    ws.freeze_panes = "A2"
    rows = "2:2000"
    for letter, options in (("D", CLASS_TITLES), ("E", ["так", "ні"]), ("H", ["I", "II", "III", "IV", "V"])):
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        dv.add(f"{letter}{rows.split(':')[0]}:{letter}{rows.split(':')[1]}")
        ws.add_data_validation(dv)

    ex = wb.create_sheet("Приклад")
    ex.append([c[0] for c in COLUMNS])
    for cell in ex[1]:
        cell.font, cell.fill = bold, fill
    for row in EXAMPLE:
        ex.append(row)
    for col, w in enumerate(widths, start=1):
        ex.column_dimensions[ex.cell(row=1, column=col).column_letter].width = w

    hs = wb.create_sheet("Як заповнювати")
    for line in HELP:
        hs.append([line])
    hs["A1"].font = Font(bold=True, size=13)
    hs.column_dimensions["A"].width = 110

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


@dataclass
class ImportResult:
    lines: list[dict] = field(default_factory=list)      # рядки чернетки
    errors: list[str] = field(default_factory=list)
    new_names: list[str] = field(default_factory=list)   # найменування, яких ще немає в довіднику


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


def _read_rows(data: bytes) -> tuple[list[tuple[int, dict]], list[str]]:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception:
        return [], ["Не вдалося прочитати файл. Потрібен файл Excel (.xlsx) — збережіть шаблон програми "
                    "у форматі «Книга Excel»."]
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [_text(h).rstrip(" *").lower() for h in next(rows, [])]
    index = {}
    for title, key, required, _tip in COLUMNS:
        if title.lower() in header:
            index[key] = header.index(title.lower())
        elif required:
            return [], [f"У першому рядку немає графи «{title}». Візьміть шаблон програми і не "
                        "змінюйте назви граф."]
    result = []
    for n, row in enumerate(rows, start=2):
        values = {key: _text(row[i]) if i < len(row) else "" for key, i in index.items()}
        if any(values.values()):
            result.append((n, values))
    wb.close()
    return result, []


def parse(conn: sqlite3.Connection, data: bytes) -> ImportResult:
    """Перевірити файл, нічого не змінюючи. Рядки — з nomenclature_id або None (нове найменування)."""
    res = ImportResult()
    rows, res.errors = _read_rows(data)
    if res.errors:
        return res
    if not rows:
        res.errors.append(f"У файлі немає жодного заповненого рядка (аркуш «{SHEET}»).")
        return res
    seen_numbers: dict[tuple[str, str], int] = {}
    new_kinds: dict[str, tuple[str, str, bool]] = {}
    for n, v in rows:
        def err(text):
            res.errors.append(f"Рядок {n}: {text}")
        if not v["name"]:
            err("не вказано найменування")
            continue
        nom_id = catalog.find_nomenclature(conn, v["name"])
        nom = conn.execute("SELECT * FROM nomenclature WHERE id = ?", (nom_id,)).fetchone() if nom_id else None
        cls_raw = v.get("class", "").lower()
        cls = CLASS_BY_TITLE.get(cls_raw, "inventory" if not cls_raw else None)
        if cls is None:
            err(f"невідомий клас обліку «{v['class']}» — оберіть зі списку: {', '.join(CLASS_TITLES)}")
            continue
        serial = (v.get("serial", "").lower() in YES) if nom is None else bool(nom["serial_tracked"])
        if nom is None:
            if not v["uom"]:
                err(f"нове найменування «{v['name']}» — вкажіть одиницю виміру")
                continue
            key = v["name"].lower()
            kind = (v["uom"], cls, serial)
            if key in new_kinds and new_kinds[key] != kind:
                err(f"«{v['name']}» в інших рядках має іншу одиницю виміру, клас або «поштучно»")
                continue
            new_kinds[key] = kind
            if v["name"] not in res.new_names:
                res.new_names.append(v["name"])
        elif v["uom"] and v["uom"] != nom["uom"]:
            err(f"«{nom['name']}» у довіднику обліковується в «{nom['uom']}», а у файлі — «{v['uom']}»")
            continue
        cat_raw = v.get("cat", "").lower().replace("кат.", "").strip()
        categorized = nom["is_categorized"] if nom is not None else 1
        cat = CATEGORIES.get(cat_raw) if cat_raw else (1 if categorized else None)
        if cat_raw and cat is None:
            err(f"невідома категорія «{v['cat']}» — I, II, III, IV або V")
            continue
        if not categorized:
            cat = None
        try:
            price = parse_money(v["price"])
            qty = parse_qty(v["qty"])
        except ValueError as exc:
            err(str(exc))
            continue
        if price is None:
            err("не вказано ціну за одиницю")
            continue
        if not qty:
            err("кількість має бути більше нуля")
            continue
        sn, inv = v.get("serial_no") or None, v.get("inventory_no") or None
        if serial:
            if not sn and not inv:
                err(f"«{v['name']}» обліковується поштучно — вкажіть заводський або інвентарний номер")
                continue
            if qty != 1000:
                err("поштучне майно — один рядок на одиницю, кількість 1")
                continue
            number = ("inv", inv) if inv else ("sn:" + v["name"].lower(), sn)
            if number in seen_numbers:
                err(f"номер {inv or sn} уже є в рядку {seen_numbers[number]}")
                continue
            seen_numbers[number] = n
        elif sn or inv:
            err(f"«{v['name']}» обліковується кількістю — номери вказують лише для поштучного "
                "(«Поштучно» — «так»)")
            continue
        res.lines.append({"row": n, "nomenclature_id": nom_id, "name": v["name"], "code": v.get("code") or None,
                          "uom": v["uom"], "class": cls, "serial": serial, "serial_no": sn, "inventory_no": inv,
                          "category": cat, "price_kop": price, "qty_m": qty, "note": v.get("note") or None})
    return res


def create_opening(conn: sqlite3.Connection, data: bytes, to_location_id: int, op_date: str) -> tuple[int, ImportResult]:
    """Перевірити файл і, якщо помилок немає, створити чернетку «Введення початкових залишків».
    Нові найменування й поштучні одиниці додаються в довідники (із журналом змін)."""
    res = parse(conn, data)
    if res.errors:
        raise DocumentError(res.errors)
    created: dict[str, int] = {}
    lines = []
    conn.execute("SAVEPOINT import_opening")   # помилка посередині — довідники не засмічуються
    try:
        for ln in res.lines:
            try:
                nom_id = ln["nomenclature_id"]
                if nom_id is None:
                    key = ln["name"].lower()
                    if key not in created:
                        created[key] = catalog.create_nomenclature(conn, ln["name"], ln["uom"], ln["class"],
                                                                   ln["serial"], code=ln["code"])
                    nom_id = created[key]
                item_id = (catalog.find_or_create_item(conn, nom_id, ln["serial_no"], ln["inventory_no"])
                           if ln["serial"] else None)
            except ValueError as exc:
                raise DocumentError([f"Рядок {ln['row']}: {exc}"]) from None
            lines.append({"nomenclature_id": nom_id, "item_id": item_id, "category": ln["category"],
                          "price_kop": ln["price_kop"], "qty_m": ln["qty_m"], "note": ln["note"]})
        doc_id = documents.create_draft(conn, "opening", locations.system(conn, locations.OPENING), to_location_id,
                                        doc_date=op_date, op_date=op_date, basis="Імпорт з Excel")
        documents.set_lines(conn, doc_id, lines)
    except Exception:
        conn.execute("ROLLBACK TO import_opening")
        conn.execute("RELEASE import_opening")
        raise
    conn.execute("RELEASE import_opening")
    return doc_id, res

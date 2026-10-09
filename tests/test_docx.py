"""Документи Word: накладна (дод. 25) за шаблоном."""
from __future__ import annotations

from io import BytesIO

from docx import Document

from oblik import docx_out
from oblik.config import Paths
from tests.test_documents_web import doc_id_from, env, open_db  # noqa: F401 — фікстура env
from tests.test_issue_owner import cable_key, post_doc, receive_cable


def text_of(data: bytes) -> str:
    d = Document(BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


def setup_issue(storage, client, ids):
    conn = open_db(storage)
    for key, value in (("unit_name", "Військова частина А0000"), ("unit_code", "А0000"),
                       ("edrpou", "12345678"), ("subunit_name", "Військова частина А1111"),
                       ("place", "м. Н-ськ"), ("operator_name", "Коваль О.")):
        conn.execute("INSERT INTO settings(key, value) VALUES (?, ?)", (key, value))
    conn.execute("UPDATE persons SET rank = 'сержант', position = 'начальник складу' WHERE id = ?", (ids["mvo"],))
    conn.execute("UPDATE persons SET rank = 'солдат' WHERE id = ?", (ids["soldier"],))
    conn.commit()
    conn.close()
    receive_cable(client, ids, f"u:{ids['wh']}", "125,5")
    return post_doc(client, "transfer", f"u:{ids['wh']}", f"p:{ids['soldier']}",
                    {"ln_key": [cable_key(ids)], "ln_qty": ["125,5"], "ln_req": ["130"], "ln_note": ["з бухти № 2"]},
                    outgoing=True)


def test_nakladna_filled(env):
    storage, client, ids = env
    doc = setup_issue(storage, client, ids)
    resp = client.get(f"/documents/{doc}/docx/nakladna")
    assert resp.status_code == 200
    assert "attachment" in resp.headers["Content-Disposition"]
    text = text_of(resp.data)
    assert "{{" not in text and "{%" not in text
    assert "(видано:" not in text   # службова позначка програми — не для паперу
    for expected in ("Військова частина А0000", "Військова частина А1111", "12345678",
                     "Накладна (вимога) № 2", "м. Н-ськ", "«02» жовтня 2026 року",
                     "Видача в користування", "Склад ТЗО", "солдат Бондар І.",
                     "Кабель", "125,5", "130", "1 255,00", "з бухти № 2",
                     "сто двадцять п'ять цілих п'ять десятих",
                     "одна тисяча двісті п'ятдесят п'ять", "Олена КОВАЛЬ",
                     "начальник складу, сержант", "Іван БОНДАР",
                     "Оброблено в електронній формі", "виконавець: Коваль О."):
        assert expected in text, expected
    # Копія — у папці output/<рік>.
    saved = list((storage.paths.output / "2026").glob("Накладна*.docx"))
    assert len(saved) == 1 and "№ 2" in saved[0].name
    # Шаблон скопійовано в templates_docx — його можна правити.
    assert (storage.paths.templates_docx / "nakladna.docx").exists()


def test_view_shows_word_button_and_draft_print(env):
    storage, client, ids = env
    doc = setup_issue(storage, client, ids)
    assert f"/documents/{doc}/docx/nakladna" in client.get(f"/documents/{doc}").get_data(as_text=True)
    # Надходження — накладної немає (буде акт).
    receipt = receive_cable(client, ids, f"u:{ids['wh']}")
    assert client.get(f"/documents/{receipt}/docx/nakladna").status_code == 302
    # Чернетка: можна роздрукувати на підпис, без відмітки про обробку.
    d = doc_id_from(client.post("/documents/new/transfer", data={
        "from_place": f"u:{ids['wh']}", "to_place": f"p:{ids['soldier']}", "doc_date": "2026-10-05"}))
    client.post(f"/documents/{d}", data={"from_place": f"u:{ids['wh']}", "to_place": f"p:{ids['soldier']}",
                                         "doc_date": "2026-10-05", "mode": "out", "action": "save",
                                         "ln_key": [cable_key(ids)], "ln_qty": ["1"], "ln_req": [""], "ln_note": [""]})
    resp = client.get(f"/documents/{d}/docx/nakladna")
    text = text_of(resp.data)
    assert "(чернетка)" in resp.headers["Content-Disposition"] or "%28" in resp.headers["Content-Disposition"]
    assert "Документ не проведено" in text


def test_user_template_and_reset(env, tmp_path):
    storage, client, ids = env
    doc = setup_issue(storage, client, ids)
    paths: Paths = storage.paths
    form = docx_out.FORMS["nakladna"]
    # Користувач змінив шаблон: додав свій рядок.
    user = docx_out.template_path(paths, form)
    d = Document(str(user))
    d.add_paragraph("Мій рядок: {{ unit_code }} / {{ count_words }}")
    d.save(str(user))
    text = text_of(client.get(f"/documents/{doc}/docx/nakladna").data)
    assert "Мій рядок: А0000 / одне" in text   # «одне найменування»
    page = client.get("/templates-docx").get_data(as_text=True)
    assert "змінений вами" in page and "processing_mark" in page
    # Зіпсований шаблон — зрозуміла помилка, а не падіння.
    d.add_paragraph("{{ зламано")
    d.save(str(user))
    resp = client.get(f"/documents/{doc}/docx/nakladna")
    assert resp.status_code == 302
    assert "Не вдалося заповнити шаблон" in client.get(resp.headers["Location"]).get_data(as_text=True)
    # Повернути стандартний; змінений — у .bak.
    resp = client.post("/templates-docx", data={"form": "nakladna"})
    assert resp.status_code == 302
    assert list(paths.templates_docx.glob("*.bak.docx"))
    assert "Мій рядок" not in text_of(client.get(f"/documents/{doc}/docx/nakladna").data)


def test_nakladna_pdf(env):
    import re
    from oblik import pdf_out
    storage, client, ids = env
    doc = setup_issue(storage, client, ids)
    resp = client.get(f"/documents/{doc}/pdf/nakladna")
    assert resp.status_code == 200 and resp.mimetype == "application/pdf"
    assert resp.data[:5] == b"%PDF-"
    assert "attachment" not in resp.headers.get("Content-Disposition", "")   # відкривається в браузері
    assert len(re.findall(rb"/Type\s*/Page[^s]", resp.data)) == 1             # одна сторінка
    assert list((storage.paths.output / "2026").glob("Накладна*.pdf"))
    # Кнопка друку на документі.
    assert f"/documents/{doc}/pdf/nakladna" in client.get(f"/documents/{doc}").get_data(as_text=True)
    # Немає шрифту — зрозуміла помилка.
    old_dirs, old_reg = pdf_out._FONT_DIRS, pdf_out._registered
    pdf_out._FONT_DIRS, pdf_out._registered = [], False
    try:
        resp = client.get(f"/documents/{doc}/pdf/nakladna")
        assert resp.status_code == 302
        assert "Не знайдено шрифт" in client.get(resp.headers["Location"]).get_data(as_text=True)
    finally:
        pdf_out._FONT_DIRS, pdf_out._registered = old_dirs, old_reg


def test_unmodified_old_standard_template_is_upgraded(env):
    storage, client, ids = env
    doc = setup_issue(storage, client, ids)
    paths = storage.paths
    form = docx_out.FORMS["nakladna"]
    user = docx_out.template_path(paths, form)
    assert docx_out.template_status(paths, form) == "standard"
    # Імітуємо оновлення програми: у папці лежить незмінена стара стандартна копія.
    old = Document(str(user))
    old.add_paragraph("стара стандартна версія")
    old.save(str(user))
    docx_out._remember(paths, form, docx_out._sha(user))
    assert docx_out.template_status(paths, form) == "old_standard"
    text = text_of(client.get(f"/documents/{doc}/docx/nakladna").data)
    assert "стара стандартна версія" not in text
    assert docx_out.template_status(paths, form) == "standard"
    # Змінений користувачем — не чіпаємо.
    mine = Document(str(user))
    mine.add_paragraph("мій підпис")
    mine.save(str(user))
    assert docx_out.template_status(paths, form) == "modified"
    assert "мій підпис" in text_of(client.get(f"/documents/{doc}/docx/nakladna").data)


def test_acts_for_receipt_split_by_class(env):
    """Надходження з ОЗ і запасами: два акти, кожен лише зі своїми рядками."""
    import re
    storage, client, ids = env
    conn = open_db(storage)
    for key, value in (("unit_name", "Військова частина А0000"), ("unit_code", "А0000"),
                       ("operator_name", "Коваль О.")):
        conn.execute("INSERT INTO settings(key, value) VALUES (?, ?)", (key, value))
    conn.commit()
    conn.close()
    doc = post_doc(client, "receipt", f"c:{ids['ext']}", f"u:{ids['wh']}", {
        "ln_nom": [str(ids["cable"]), str(ids["cam"])], "ln_name": ["Кабель", "Камера"],
        "ln_serial": ["", "SN-555"], "ln_inv": ["", ""], "ln_cat": ["1", "2"], "ln_price": ["10", "4460"],
        "ln_qty": ["30", ""], "ln_req": ["", ""], "ln_note": ["", ""]}, outgoing=False, op_date="2026-10-01")
    conn = open_db(storage)
    assert [f.code for f in docx_out.available_forms(conn, doc)] == ["act_oz", "act_zap"]
    conn.close()
    oz = text_of(client.get(f"/documents/{doc}/docx/act_oz").data)
    zap = text_of(client.get(f"/documents/{doc}/docx/act_zap").data)
    for text in (oz, zap):
        assert "{{" not in text and "{%" not in text and "ЗАТВЕРДЖУЮ" in text and "Вища частина" in text
    assert "Камера" in oz and "SN-555" in oz and "Кабель" not in oz and "4 460,00" in oz
    assert "Кабель" in zap and "Камера" not in zap and "тридцять одиниць" in zap and "триста грн" in zap
    assert "приймає" in zap   # МВО нашої сторони приймає на відповідальне зберігання
    for code in ("act_oz", "act_zap"):
        resp = client.get(f"/documents/{doc}/pdf/{code}")
        assert resp.status_code == 200 and resp.data[:5] == b"%PDF-"
    view = client.get(f"/documents/{doc}").get_data(as_text=True)
    assert f"/documents/{doc}/pdf/act_oz" in view and f"/documents/{doc}/pdf/act_zap" in view
    assert "nakladna" not in view   # для надходження накладної немає
    # Лише запаси — акта ОЗ не пропонує.
    only = receive_cable(client, ids, f"u:{ids['wh']}")
    conn = open_db(storage)
    assert [f.code for f in docx_out.available_forms(conn, only)] == ["act_zap"]
    conn.close()
    assert client.get(f"/documents/{only}/pdf/act_oz").status_code == 302
    assert len(re.findall(r"act_", client.get("/templates-docx").get_data(as_text=True))) >= 2

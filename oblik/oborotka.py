"""Імпорт оборотки з бухгалтерії (Excel) → чернетки «Введення початкових залишків».

Оборотка — плоска таблиця: Рахунок | Контрагент (МВО) | ПБО | Найменування номенклатури |
Основна одиниця виміру | Номенклатура | Ціна | Кількість | Сума (назви граф бувають трохи
різні — «Нова ціна (бухгалтерська оцінка)», «Кількість на кінець», «Залишок на кінець»).
У файлі — майно кількох МВО; користувач обирає, чиє імпортувати і на яке місце обліку.

Клас майна і поштучний облік — за рахунком (типові значення — припущення за планом рахунків
бюджетних установ, користувач бачить і змінює їх перед імпортом). Поштучне майно: інвентарний
номер = графа «Номенклатура». Категорія — I (у оборотці її немає; виправляється в чернетці).

© 2026 Yevhenii Hosting by LLC Hosting in Ukraine
"""
from __future__ import annotations

import re
import sqlite3
from collections import OrderedDict
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from io import BytesIO

from . import catalog, documents, locations
from .documents import DocumentError
from .textutil import QTY_SCALE, format_money

# «(1450,00 - 30.07.2026)» наприкінці назви — первісна ціна й дата з бухгалтерії (бувають і
# з помилками на кшталт «(1,00-21.12.2023}»). Це не частина назви: інакше те саме найменування
# після ручного виправлення назви стало б дублем.
_PRICE_DATE_TAIL = re.compile(r"\s*[({]\s*\d[\d\s.,]*[-–—]\s*\d{1,2}\.\d{1,2}\.\d{2,4}\s*[)}]\s*$")


def clean_name(name: str) -> str:
    return _PRICE_DATE_TAIL.sub("", name).strip()


# Графа → варіанти заголовка (нижній регістр, без зайвих пробілів). Перевіряються по черзі:
# точний збіг або початок заголовка.
HEADERS = OrderedDict([
    ("account", ["рахунок", "субрахунок", "счет"]),
    ("mvo", ["контрагент", "мво", "матеріально відповідальна особа"]),
    ("name", ["найменування номенклатури", "найменування", "наименование номенклатуры", "наименование"]),
    ("uom", ["основна одиниця виміру", "одиниця виміру", "од. виміру", "единица измерения"]),
    ("code", ["номенклатура", "код номенклатури"]),
    ("price", ["нова ціна", "ціна", "цена"]),
    ("qty", ["кількість", "количество"]),
    ("sum", ["сума", "залишок на кінець", "сумма"]),
])
REQUIRED = ("account", "mvo", "name", "uom", "price", "qty")

# Типовий клас за початком рахунку: (клас, поштучно, якщо скрізь кількість 1).
DEFAULT_CLASS = [("101", "fixed", True), ("131", "fixed", True), ("111", "low_value", False),
                 ("15", "inventory", False), ("18", "mshp", False)]

CENT = Decimal("0.01")


def _norm(v) -> str:
    return " ".join(str(v or "").split()).lower()


def _header_index(header: list) -> dict[str, int] | None:
    names = [_norm(h) for h in header]
    index: dict[str, int] = {}
    for key, variants in HEADERS.items():
        for v in variants:
            hit = next((i for i, n in enumerate(names)
                        if i not in index.values() and (n == v or n.startswith(v + " ") or n.startswith(v + " ("))),
                       None)
            if hit is not None:
                index[key] = hit
                break
    return index if all(k in index for k in REQUIRED) else None


def is_oborotka(data: bytes) -> bool:
    """Чи схожий файл на оборотку (за заголовками першого рядка)."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
        header = next(wb.worksheets[0].iter_rows(max_row=1, values_only=True), [])
        wb.close()
    except Exception:
        return False
    return _header_index(list(header)) is not None


def _dec(v) -> Decimal | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        return Decimal(str(v).replace(" ", "").replace(" ", "").replace(",", "."))
    except InvalidOperation:
        return None


@dataclass
class Row:
    n: int                 # рядок Excel
    account: str
    mvo: str
    name: str
    uom: str
    code: str
    price: Decimal | None
    qty: Decimal | None
    sum: Decimal | None


@dataclass
class Account:
    code: str
    rows: int = 0
    all_one: bool = True          # у всіх рядках кількість 1 — можна вести поштучно
    accounting_class: str = "inventory"
    serial: bool = False


@dataclass
class Mvo:
    name: str
    rows: int = 0
    total: Decimal = Decimal(0)

    @property
    def total_text(self) -> str:
        return format_money(int((self.total * 100).quantize(Decimal(1), ROUND_HALF_UP)))


@dataclass
class Oborotka:
    rows: list[Row] = field(default_factory=list)
    mvos: list[Mvo] = field(default_factory=list)
    accounts: list[Account] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def read(data: bytes) -> Oborotka:
    """Прочитати оборотку: рядки, МВО (з кількістю рядків і сумою), рахунки з типовим класом."""
    from openpyxl import load_workbook
    res = Oborotka()
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception:
        res.errors.append("Не вдалося прочитати файл. Потрібен файл Excel (.xlsx).")
        return res
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    index = _header_index(list(next(it, [])))
    if index is None:
        res.errors.append("Це не схоже на оборотку: у першому рядку мають бути графи «Рахунок», "
                          "«Контрагент», «Найменування номенклатури», «Одиниця виміру», «Ціна», «Кількість».")
        wb.close()
        return res
    mvos: dict[str, Mvo] = {}
    accounts: dict[str, Account] = {}
    for n, row in enumerate(it, start=2):
        get = lambda k: row[index[k]] if k in index and index[k] < len(row) else None  # noqa: E731
        name = clean_name(" ".join(str(get("name") or "").split()))
        mvo = " ".join(str(get("mvo") or "").split())
        if not name and not mvo:
            continue
        low = _norm(get("account"))
        if low.startswith(("разом", "всього", "усього", "итого")):
            continue
        r = Row(n, " ".join(str(get("account") or "").split()), mvo, name,
                " ".join(str(get("uom") or "").split()), " ".join(str(get("code") or "").split()),
                _dec(get("price")), _dec(get("qty")), _dec(get("sum")))
        res.rows.append(r)
        m = mvos.setdefault(mvo, Mvo(mvo))
        m.rows += 1
        m.total += r.sum if r.sum is not None else (r.price or 0) * (r.qty or 0)
        a = accounts.setdefault(r.account, Account(r.account))
        a.rows += 1
        a.all_one = a.all_one and r.qty == 1
    wb.close()
    for a in accounts.values():
        for prefix, cls, serial in DEFAULT_CLASS:
            if a.code.startswith(prefix):
                a.accounting_class, a.serial = cls, serial and a.all_one
                break
    res.mvos = sorted(mvos.values(), key=lambda m: m.name.lower())
    res.accounts = sorted(accounts.values(), key=lambda a: a.code)
    return res


@dataclass
class Plan:
    places: dict[str, int]                       # МВО → location_id (лише обрані)
    classes: dict[str, tuple[str, bool]]         # рахунок → (клас, поштучно)
    op_date: str
    categories: dict[str, int] = field(default_factory=dict)   # рахунок → категорія (I–V = 1–5), типово 1


@dataclass
class Result:
    doc_ids: list[int] = field(default_factory=list)
    lines: int = 0
    new_names: int = 0
    rounding: Decimal = Decimal(0)               # розбіжність з «Сума» через округлення ціни до копійок
    rounding_rows: int = 0


def _line(conn, r: Row, cls: str, serial: bool, created: dict, res: Result, category: int = 1) -> dict:
    if not r.name:
        raise ValueError("не вказано найменування")
    if r.price is None or r.price < 0:
        raise ValueError("ціна не вказана або від'ємна")
    if r.qty is None or r.qty <= 0:
        raise ValueError("кількість не вказана або не більше нуля")
    qty_m = r.qty * QTY_SCALE
    if qty_m != qty_m.to_integral_value():
        raise ValueError("кількість — не більше трьох знаків після коми")
    if serial and r.qty != 1:
        raise ValueError("поштучне майно (рахунок обліковується поштучно) — кількість має бути 1")
    if serial and not r.code:
        raise ValueError("поштучне майно — немає номера в графі «Номенклатура» (інвентарний номер)")
    price_kop = int((r.price * 100).quantize(Decimal(1), ROUND_HALF_UP))
    if r.sum is not None:
        diff = abs(price_kop * r.qty / 100 - r.sum)
        if diff >= CENT:
            res.rounding += diff
            res.rounding_rows += 1
    nom_id = catalog.find_nomenclature(conn, r.name, None if serial else r.code)
    if nom_id is not None:
        nom = conn.execute("SELECT * FROM nomenclature WHERE id = ?", (nom_id,)).fetchone()
        if nom["uom"] != r.uom and r.uom:
            raise ValueError(f"«{nom['name']}» у довіднику обліковується в «{nom['uom']}», в оборотці — «{r.uom}»")
        if bool(nom["serial_tracked"]) != serial:
            raise ValueError(f"«{nom['name']}» у довіднику " + ("поштучне" if nom["serial_tracked"] else "кількісне")
                             + " — а рахунок " + r.account + (" обліковується поштучно" if serial else " — кількістю"))
    else:
        key = r.name.lower()
        if key not in created:
            created[key] = catalog.create_nomenclature(conn, r.name, r.uom or "шт.", cls, serial,
                                                       code=None if serial else (r.code or None),
                                                       account=r.account or None)
            res.new_names += 1
        nom_id = created[key]
    item_id = catalog.find_or_create_item(conn, nom_id, None, r.code) if serial else None
    nom = conn.execute("SELECT is_categorized FROM nomenclature WHERE id = ?", (nom_id,)).fetchone()
    return {"nomenclature_id": nom_id, "item_id": item_id, "category": category if nom["is_categorized"] else None,
            "price_kop": price_kop, "qty_m": int(qty_m)}


def create_drafts(conn: sqlite3.Connection, data: bytes, plan: Plan) -> Result:
    """Чернетки «Введення початкових залишків» — по одній на кожне обране місце обліку.
    Якщо хоч один рядок з помилкою — нічого не створюється (DocumentError з усіма помилками)."""
    ob = read(data)
    if ob.errors:
        raise DocumentError(ob.errors)
    if not plan.places:
        raise DocumentError("Позначте хоча б одну МВО, чиє майно імпортувати")
    res = Result()
    errors: list[str] = []
    by_place: dict[int, list[dict]] = OrderedDict()
    created: dict[str, int] = {}
    conn.execute("SAVEPOINT import_oborotka")
    try:
        for r in ob.rows:
            if r.mvo not in plan.places:
                continue
            cls, serial = plan.classes.get(r.account, ("inventory", False))
            try:
                ln = _line(conn, r, cls, serial, created, res, plan.categories.get(r.account, 1))
            except ValueError as exc:
                errors.append(f"Рядок {r.n}: {exc}")
                continue
            by_place.setdefault(plan.places[r.mvo], []).append(ln)
        seen_items: dict[int, int] = {}
        for lines in by_place.values():
            for ln in lines:
                if ln["item_id"] is not None:
                    seen_items[ln["item_id"]] = seen_items.get(ln["item_id"], 0) + 1
        if any(n > 1 for n in seen_items.values()):
            errors.append("У файлі той самий інвентарний номер у кількох рядках обраних МВО")
        if errors:
            raise DocumentError(errors)
        opening = locations.system(conn, locations.OPENING)
        for place, lines in by_place.items():
            doc_id = documents.create_draft(conn, "opening", opening, place, doc_date=plan.op_date,
                                            op_date=plan.op_date, basis="Оборотка з бухгалтерії")
            documents.set_lines(conn, doc_id, lines)
            res.doc_ids.append(doc_id)
            res.lines += len(lines)
    except Exception:
        conn.execute("ROLLBACK TO import_oborotka")
        conn.execute("RELEASE import_oborotka")
        raise
    conn.execute("RELEASE import_oborotka")
    return res

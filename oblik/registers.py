"""Облікові регістри за формами Інструкції (наказ МОУ № 440):

  * узагальнююча відомість обліку військового майна в електронній формі — дод. 1
    (п. 7 розд. II): підсумки надходження, вибуття і наявності за звітний період;
  * книга реєстрації та руху облікових документів — дод. 2 (п. 10 розд. II).

Функції нічого не змінюють у базі.
"""
from __future__ import annotations

import sqlite3
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date
from io import BytesIO

from . import balances, documents, locations, refs, settings
from .documents import DOC_TYPES
from .reports import ReportError, resolve_scope
from .textutil import QTY_SCALE, person_signature, uk_sort_key

FIXED_CLASSES = {"fixed", "low_value"}


# --- Підписанти -------------------------------------------------------------------

def signer_line(conn: sqlite3.Connection, person_id: int | None) -> str:
    """«посада, військове звання, ____ Власне ім'я ПРІЗВИЩЕ» — як у підписах форм."""
    if not person_id:
        return ""
    p = conn.execute("SELECT * FROM persons WHERE id = ?", (person_id,)).fetchone()
    if p is None:
        return ""
    parts = [x for x in (p["position"], p["rank"]) if x]
    head = ", ".join(parts)
    return f"{head}{', ' if head else ''}____________ {person_signature(p['last_name'], p['first_name'])}"


def _scope_mvo(conn, scope: str) -> int | None:
    """МВО для підпису «веде облік» за вибором «чиї»."""
    kind, _, raw = scope.partition(":")
    if not raw.isdigit():
        return None
    if kind == "m":
        return int(raw)
    if kind == "u":
        row = conn.execute("SELECT mvo_person_id FROM units WHERE id = ?", (int(raw),)).fetchone()
        return row["mvo_person_id"] if row else None
    return None


# --- Дод. 1. Узагальнююча відомість -----------------------------------------------

@dataclass
class Summary:
    title: str
    date_from: str
    date_to: str
    rows: list[dict] = field(default_factory=list)
    default_book: str = ""
    default_keeper_id: int | None = None
    default_responsible_id: int | None = None


def summary(conn: sqlite3.Connection, scope: str, date_from: str, date_to: str,
            include_subunits: bool = False) -> Summary:
    """Надійшло / вибуло за період і наявне на кінець — по кожному найменуванню.

    Рух усередині обраних місць (наприклад, видача зі служби військовослужбовцю того ж
    підрозділу) не є ні надходженням, ні вибуттям — майно лишається в тій самій книзі.
    Сторно не збільшує обороти, а зменшує те, що скасовує.
    """
    if date_from > date_to:
        raise ReportError("Початок звітного періоду пізніше за кінець")
    location_ids, title = resolve_scope(conn, scope, include_subunits)
    if location_ids is None:
        scope_set = {r["id"] for r in conn.execute(
            "SELECT id FROM locations WHERE kind IN ('unit', 'person')")}
    else:
        scope_set = set(location_ids)

    incoming: dict[int, int] = defaultdict(int)
    outgoing: dict[int, int] = defaultdict(int)
    for m in conn.execute(
            "SELECT m.from_location_id, m.to_location_id, m.nomenclature_id, m.qty_m, d.doc_type "
            "FROM movements m JOIN documents d ON d.id = m.document_id "
            "WHERE m.op_date BETWEEN ? AND ?", (date_from, date_to)):
        src_in, dst_in = m["from_location_id"] in scope_set, m["to_location_id"] in scope_set
        if src_in == dst_in:
            continue
        nom, q, storno = m["nomenclature_id"], m["qty_m"], m["doc_type"] == "storno"
        if dst_in:
            if storno:
                outgoing[nom] -= q   # повернення вибулого назад = скасування вибуття
            else:
                incoming[nom] += q
        else:
            if storno:
                incoming[nom] -= q   # скасування надходження
            else:
                outgoing[nom] += q

    end: dict[int, int] = defaultdict(int)
    for b in balances.balances(conn, sorted(scope_set), date_to):
        end[b["nomenclature_id"]] += b["qty_m"]

    noms = {r["id"]: r for r in conn.execute("SELECT id, name, uom FROM nomenclature")}
    report = Summary(title=title, date_from=date_from, date_to=date_to)
    for nom_id in set(incoming) | set(outgoing) | set(end):
        if not (incoming[nom_id] or outgoing[nom_id] or end[nom_id]):
            continue
        n = noms[nom_id]
        report.rows.append({"nomenclature_id": nom_id, "name": n["name"], "uom": n["uom"],
                            "in_m": incoming[nom_id], "out_m": outgoing[nom_id], "end_m": end[nom_id]})
    report.rows.sort(key=lambda r: uk_sort_key(r["name"]))
    for i, r in enumerate(report.rows, 1):
        r["n"] = i

    s = settings.load(conn)
    report.default_book = "Книга обліку військового майна — " + title
    report.default_keeper_id = _scope_mvo(conn, scope)
    report.default_responsible_id = s.get("service_head_person_id")
    return report


# --- Дод. 2. Книга реєстрації та руху облікових документів ---------------------------

def paper_title(conn: sqlite3.Connection, doc) -> str:
    """Найменування облікового документа для графи 3."""
    t = doc["doc_type"]
    if t in ("receipt", "opening"):
        classes = {r[0] for r in conn.execute(
            "SELECT DISTINCT n.accounting_class FROM document_lines l "
            "JOIN nomenclature n ON n.id = l.nomenclature_id WHERE l.document_id = ?", (doc["id"],))}
        if t == "opening":
            return "Введення початкових залишків"
        if classes and classes <= FIXED_CLASSES:
            return "Акт приймання-передачі основних засобів"
        if classes and not classes & FIXED_CLASSES:
            return "Акт приймання-передачі запасів"
        return "Акт приймання-передачі основних засобів / запасів"
    if t in ("transfer", "dispatch"):
        return "Накладна (вимога)"
    if t == "writeoff":
        return "Акт списання"
    if t == "order":
        return "Наряд на видавання (приймання) військового майна"
    if t == "condition":
        return "Акт якісного (технічного) стану"
    if t == "storno":
        orig = documents.get(conn, doc["reversal_of_id"])
        return f"Сторно документа рег. № {orig['reg_no']}"
    return DOC_TYPES[t].title


def _counterpart(conn, doc) -> str:
    """Графа 9: від кого надійшов або кому переданий документ."""
    src = locations.title(conn, doc["from_location_id"])
    dst = locations.title(conn, doc["to_location_id"])
    t = doc["doc_type"]
    if t == "receipt":
        return src
    if t in ("dispatch", "opening"):
        return dst
    if t in ("writeoff", "condition"):
        return src
    return f"{src} → {dst}"


def _goods(conn, doc_id: int, limit: int = 4) -> str:
    """Графа 8: на яке майно документ."""
    names = [r[0] for r in conn.execute(
        "SELECT n.name FROM document_lines l JOIN nomenclature n ON n.id = l.nomenclature_id "
        "WHERE l.document_id = ? GROUP BY n.id ORDER BY MIN(l.line_no)", (doc_id,))]
    if len(names) > limit:
        return ", ".join(names[:limit]) + f" та ще {len(names) - limit}"
    return ", ".join(names)


@dataclass
class Journal:
    year: int
    rows: list[dict] = field(default_factory=list)
    drafts: int = 0


def journal(conn: sqlite3.Connection, year: int) -> Journal:
    """Зареєстровані документи року за порядком реєстраційних номерів.
    Чернетки не реєструються (документ реєструють після підписання), анульовані — показуються."""
    result = Journal(year=year)
    result.drafts = conn.execute("SELECT COUNT(*) FROM documents WHERE reg_year = ? AND status = 'draft'",
                                 (year,)).fetchone()[0]
    for d in conn.execute("SELECT * FROM documents WHERE reg_year = ? AND status <> 'draft' "
                          "ORDER BY reg_no", (year,)):
        cancelled = d["status"] == "cancelled"
        result.rows.append({
            "id": d["id"],
            "reg_no": d["reg_no"],
            "reg_date": documents.fmt_date((d["posted_at"] or d["doc_date"])[:10]),
            "title": paper_title(conn, d) + (" — анульовано" if cancelled else ""),
            "doc_no": d["doc_no"],
            "doc_date": documents.fmt_date(d["doc_date"]),
            "goods": "" if cancelled else _goods(conn, d["id"]),
            "counterpart": "" if cancelled else _counterpart(conn, d),
            "valid_until": documents.fmt_date(d["valid_until"]),
            "cancelled": cancelled,
            "reversed": bool(d["reversed_by_id"]),
        })
    return result


# --- Excel ------------------------------------------------------------------------

def _xl():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter
    thin = Side(style="thin")
    return Workbook, Alignment, Font, get_column_letter, Border(left=thin, right=thin, top=thin, bottom=thin)


def _print_setup(ws, landscape: bool) -> None:
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _qty(milli: int) -> float:
    return milli / QTY_SCALE


def summary_xlsx(conn: sqlite3.Connection, report: Summary, number: str, book: str,
                 book_details: str, compiled: str, responsible_id: int | None,
                 keeper_id: int | None) -> bytes:
    Workbook, Alignment, Font, letter, border = _xl()
    s = settings.load(conn)
    wb = Workbook()
    ws = wb.active
    ws.title = "Дод. 1"
    widths = [7, 52, 11, 16, 16, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[letter(i)].width = w
    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(wrap_text=True, horizontal="center", vertical="center")

    def line(text, bold=False, size=None, align=None):
        ws.append([text])
        r = ws.max_row
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        c = ws.cell(r, 1)
        c.font = Font(bold=bold, size=size)
        c.alignment = align or wrap

    for name in (s.get("unit_name"), s.get("subunit_name")):
        if name:
            line(name)
    line(f"Узагальнююча відомість № {number or '______'}", bold=True, size=13, align=center)
    line("обліку військового майна в електронній формі", bold=True, align=center)
    ws.append([])
    period = f"{documents.fmt_date(report.date_from)} – {documents.fmt_date(report.date_to)}"
    for label, value in (("Найменування книги обліку військового майна", book),
                         ("Реквізити книги обліку військового майна", book_details),
                         ("Звітний період", period),
                         ("Дата складання узагальнюючої відомості", documents.fmt_date(compiled))):
        ws.append([label, None, None, value])
        r = ws.max_row
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
        for col in (1, 4):
            ws.cell(r, col).alignment = wrap
            ws.cell(r, col).border = border
        for col in range(1, 7):
            ws.cell(r, col).border = border
    ws.append([])

    heads = ["№ з/п", "Найменування військового майна", "Одиниця виміру",
             "Кількість військового майна, що надійшло", "Кількість військового майна, що вибуло",
             "Кількість наявного військового майна"]
    ws.append(heads)
    head_row = ws.max_row
    ws.append([1, 2, 3, 4, 5, 6])
    for r in (head_row, head_row + 1):
        for c in ws[r]:
            c.alignment = center
            c.border = border
            c.font = Font(bold=r == head_row)
    for r in report.rows:
        ws.append([r["n"], r["name"], r["uom"], _qty(r["in_m"]), _qty(r["out_m"]), _qty(r["end_m"])])
        for c in ws[ws.max_row]:
            c.border = border
            c.alignment = wrap
        for col in (4, 5, 6):
            ws.cell(ws.max_row, col).number_format = "# ##0.###"
    ws.print_title_rows = f"{head_row}:{head_row + 1}"

    for person_id, caption in (
            (responsible_id, "(посада, військове звання, підпис, власне ім'я, прізвище посадової (службової) "
                             "особи, яка відповідає за стан обліку військового майна, або проводить звірку)"),
            (keeper_id, "(посада, військове звання, підпис, власне ім'я, прізвище посадової (службової) "
                        "особи, яка безпосередньо веде його облік, або з якою проводиться звірка)")):
        ws.append([])
        line(signer_line(conn, person_id) or "_" * 90)
        line(caption, size=8)
    _print_setup(ws, landscape=False)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def journal_xlsx(conn: sqlite3.Connection, j: Journal) -> bytes:
    Workbook, Alignment, Font, letter, border = _xl()
    s = settings.load(conn)
    labels = refs.Labels(conn)
    wb = Workbook()
    ws = wb.active
    ws.title = f"Книга реєстрації {j.year}"
    widths = [9, 11, 30, 10, 11, 10, 10, 34, 30, 11, 14, 14, 9, 9]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[letter(i)].width = w
    center = Alignment(wrap_text=True, horizontal="center", vertical="center")
    wrap = Alignment(wrap_text=True, vertical="top")

    def line(text, bold=False, size=None):
        ws.append([text])
        r = ws.max_row
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=14)
        ws.cell(r, 1).font = Font(bold=bold, size=size)
        ws.cell(r, 1).alignment = center

    line("КНИГА", bold=True, size=13)
    line("реєстрації та руху облікових документів", bold=True)
    service = labels.get("services", s.get("default_service_id"))
    line(service or "____________________ (служба забезпечення)")
    line(", ".join(x for x in (s.get("subunit_name"), s.get("unit_name")) if x)
         or "____________________ (військова частина)")
    line(f"{j.year} рік")
    ws.append([])

    top = ["Реєстраційний номер", "Дата реєстрації", "Найменування документа", "Номер документа",
           "Дата документа", "Кількість примірників", "Загальна кількість аркушів",
           "Військове майно, на яке видано (надійшов) документ",
           "Від кого надійшов або кому переданий документ на виконання",
           "Строк виконання документа (дата)", "Підпис про одержання документа і дата",
           "Підпис про прийняття виконаного документа і дата",
           "Місцезнаходження виконаного документа", None]
    ws.append(top)
    r1 = ws.max_row
    ws.append([None] * 12 + ["номер справи", "номери аркушів у справі"])
    r2 = ws.max_row
    ws.append(list(range(1, 15)))
    r3 = ws.max_row
    for col in range(1, 13):
        ws.merge_cells(start_row=r1, start_column=col, end_row=r2, end_column=col)
    ws.merge_cells(start_row=r1, start_column=13, end_row=r1, end_column=14)
    # Excel не підбирає висоту для об'єднаних клітинок — задаємо самі.
    ws.row_dimensions[r1].height = 45
    ws.row_dimensions[r2].height = 45
    for r in (r1, r2, r3):
        for col in range(1, 15):
            c = ws.cell(r, col)
            c.alignment = center
            c.border = border
            c.font = Font(bold=r != r3, size=9)
    for row in j.rows:
        ws.append([row["reg_no"], row["reg_date"], row["title"], row["doc_no"], row["doc_date"],
                   None, None, row["goods"], row["counterpart"], row["valid_until"],
                   None, None, None, None])
        for c in ws[ws.max_row]:
            c.border = border
            c.alignment = wrap
    ws.print_title_rows = f"{r1}:{r3}"
    _print_setup(ws, landscape=True)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def default_period(today: date | None = None) -> tuple[str, str]:
    """Звітний період за замовчуванням — поточний місяць до сьогодні."""
    today = today or date.today()
    return today.replace(day=1).isoformat(), today.isoformat()

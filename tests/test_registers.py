"""Узагальнююча відомість (дод. 1) і книга реєстрації документів (дод. 2)."""
from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from oblik import db, documents, registers
from oblik.config import Paths
from oblik.reports import ReportError
from oblik.storage import Storage
from oblik.web import create_app
from tests import test_posting
from tests.test_posting import World


def cam(world, item):
    return {"nomenclature_id": world.camera, "item_id": item, "qty_m": 1000, "price_kop": 446000,
            "category": 1}


def build(world):
    """Жовтень: надходження, видача особі, передача в роту, передача назовні + її сторно, списання."""
    cable = test_posting.cable
    ids = {}
    ids["receipt"] = world.doc("receipt", world.ext, world.wh,
                               [cable(100), cam(world, world.cam1), cam(world, world.cam2)],
                               op_date="2026-10-01")
    ids["issue"] = world.doc("transfer", world.wh, world.person, [cam(world, world.cam1)], op_date="2026-10-05")
    ids["to_coy"] = world.doc("transfer", world.wh, world.coy, [cable(30)], op_date="2026-10-06")
    ids["dispatch"] = world.doc("dispatch", world.wh, world.ext, [cable(10)], op_date="2026-10-10")
    ids["storno"] = documents.reverse(world.conn, ids["dispatch"], on_date="2026-10-12")
    ids["writeoff"] = world.doc("writeoff", world.coy, world.written_off, [cable(5)], op_date="2026-10-15")
    return ids


@pytest.fixture
def w(tmp_path):
    conn = db.connect(tmp_path / "t.db")
    db.migrate(conn)
    db.init_meta(conn, "test")
    world = World(conn)
    test_posting.W = world
    world.ids = build(world)
    yield world
    conn.close()


def table(report):
    return {r["name"]: (r["in_m"] // 1000, r["out_m"] // 1000, r["end_m"] // 1000) for r in report.rows}


def test_summary_for_warehouse(w):
    rep = registers.summary(w.conn, f"u:{w.wh_unit}", "2026-10-01", "2026-10-31")
    # Кабель: 100 надійшло; вибуло 30 у роту (передача назовні сторнована — не рахується).
    # Камера: 2 надійшли, 1 видана бійцю (він не зі складу — вибуття).
    assert table(rep) == {"Кабель": (100, 30, 70), "Камера": (2, 1, 1)}
    assert [r["n"] for r in rep.rows] == [1, 2]
    assert rep.default_keeper_id == w.mvo
    # Частина періоду: лише рух після 06.10, наявне — на кінець.
    rep = registers.summary(w.conn, f"u:{w.wh_unit}", "2026-10-06", "2026-10-31")
    assert table(rep) == {"Кабель": (0, 30, 70), "Камера": (0, 0, 1)}


def test_summary_whole_unit_ignores_internal_moves(w):
    rep = registers.summary(w.conn, "all", "2026-10-01", "2026-10-31")
    # Усередині частини рух не рахується; вибуло лише списане.
    assert table(rep) == {"Кабель": (100, 5, 95), "Камера": (2, 0, 2)}
    # Рота з виданим бійцю: камера надійшла (видана бійцю роти), кабель 30 − 5.
    rep = registers.summary(w.conn, f"u:{w.coy_unit}", "2026-10-01", "2026-10-31")
    assert table(rep) == {"Кабель": (30, 5, 25), "Камера": (1, 0, 1)}


def test_summary_errors_and_empty(w):
    with pytest.raises(ReportError, match="пізніше"):
        registers.summary(w.conn, "all", "2026-10-31", "2026-10-01")
    assert registers.summary(w.conn, "all", "2026-09-01", "2026-09-30").rows == []


def test_summary_xlsx(w):
    w.conn.execute("UPDATE persons SET position = 'начальник складу', rank = 'сержант' WHERE id = ?",
                   (w.mvo,))
    rep = registers.summary(w.conn, f"u:{w.wh_unit}", "2026-10-01", "2026-10-31")
    data = registers.summary_xlsx(w.conn, rep, "7", rep.default_book, "рапорт № 1", "2026-11-01",
                                  None, w.mvo)
    flat = [v for row in load_workbook(BytesIO(data)).active.iter_rows(values_only=True) for v in row if v]
    assert "Узагальнююча відомість № 7" in flat and "01.10.2026 – 31.10.2026" in flat
    assert "Кількість наявного військового майна" in flat and 70 in flat
    assert any(isinstance(v, str) and v.startswith("начальник складу, сержант") and "Мво СКЛАД" in v
               for v in flat)


def test_journal(w):
    # Анульована чернетка і звичайна чернетка.
    cancelled = w.doc("transfer", w.wh, w.coy, [test_posting.cable(1)], post=False)
    documents.cancel(w.conn, cancelled)
    w.doc("transfer", w.wh, w.coy, [test_posting.cable(1)], post=False)
    j = registers.journal(w.conn, 2026)
    assert j.drafts == 1
    rows = {r["reg_no"]: r for r in j.rows}
    assert [r["reg_no"] for r in j.rows] == sorted(rows)
    receipt = documents.get(w.conn, w.ids["receipt"])
    r = rows[receipt["reg_no"]]
    assert r["title"] == "Акт приймання-передачі основних засобів / запасів"
    assert r["goods"] == "Кабель, Камера" and r["counterpart"] == "Вища частина"
    assert rows[documents.get(w.conn, w.ids["issue"])["reg_no"]]["title"] == "Накладна (вимога)"
    dispatch = documents.get(w.conn, w.ids["dispatch"])
    assert rows[dispatch["reg_no"]]["reversed"]
    storno = rows[documents.get(w.conn, w.ids["storno"])["reg_no"]]
    assert storno["title"] == f"Сторно документа рег. № {dispatch['reg_no']}"
    c = rows[documents.get(w.conn, cancelled)["reg_no"]]
    assert c["cancelled"] and c["title"].endswith("анульовано") and c["goods"] == ""

    ws = load_workbook(BytesIO(registers.journal_xlsx(w.conn, j))).active
    flat = [v for row in ws.iter_rows(values_only=True) for v in row if v]
    assert "реєстрації та руху облікових документів" in flat and "номер справи" in flat
    assert 14 in flat and "Накладна (вимога)" in flat


def test_pages(tmp_path):
    storage = Storage(Paths(root=tmp_path / "Флешка", local=tmp_path / "ПК"))
    storage.open()
    conn = db.connect(storage.work_db)
    world = World(conn)
    test_posting.W = world
    build(world)
    conn.commit()
    conn.close()
    app = create_app(storage)
    app.config["TESTING"] = True
    client = app.test_client()
    html = client.get("/reports").get_data(as_text=True)
    assert "/reports/summary" in html and "/reports/journal" in html
    url = f"/reports/summary?scope=u:{world.wh_unit}&from=2026-10-01&to=2026-10-31"
    html = client.get(url).get_data(as_text=True)
    assert "Кабель" in html and "Зберегти в Excel" in html
    assert "summary.xlsx" in html
    resp = client.get(url.replace("/summary?", "/summary.xlsx?") + "&number=3")
    assert resp.status_code == 200 and resp.data[:2] == b"PK"
    assert "пізніше" in client.get("/reports/summary?from=2026-10-31&to=2026-10-01").get_data(as_text=True)
    assert client.get("/reports/summary").status_code == 200
    html = client.get("/reports/journal?year=2026").get_data(as_text=True)
    assert "Накладна (вимога)" in html
    resp = client.get("/reports/journal.xlsx?year=2026")
    assert resp.status_code == 200 and resp.data[:2] == b"PK"

"""Створює стандартні Word-шаблони документів (docxtpl) у oblik/docx_templates/.

Запуск: .venv\\Scripts\\python tools\\make_docx_templates.py
Шаблони — за формами додатків до Інструкції (наказ МОУ № 440). Поля — {{ назва }},
рядки таблиці — {%tr for l in lines %} … {%tr endfor %}. Перелік полів — сторінка
«Документи Word» у програмі (oblik/docx_out.py, FIELDS).
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUT = Path(__file__).resolve().parent.parent / "oblik" / "docx_templates"
FONT = "Times New Roman"


def new_doc(landscape: bool = False) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.left_margin, sec.right_margin = Cm(2), Cm(1)
    sec.top_margin, sec.bottom_margin = Cm(1.5), Cm(1.5)
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(11)  # як у PDF (pdf_out) — щоб уміщалося на аркуш
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = style.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    compat = OxmlElement("w:compat")
    setting = OxmlElement("w:compatSetting")
    for k, v in (("w:name", "compatibilityMode"), ("w:uri", "http://schemas.microsoft.com/office/word"),
                 ("w:val", "15")):
        setting.set(qn(k), v)
    compat.append(setting)
    old = doc.settings.element.find(qn("w:compat"))
    if old is not None:
        doc.settings.element.remove(old)
    doc.settings.element.append(compat)
    return doc


def para(doc_or_cell, text: str = "", size: float | None = None, bold: bool = False,
         align=None, italic: bool = False, space_after: float = 0):
    p = doc_or_cell.add_paragraph()
    if text:
        run = p.add_run(text)
        run.bold, run.italic = bold, italic
        if size:
            run.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def gap(doc, pt: float) -> None:
    """Невеликий вертикальний проміжок (порожній абзац висотою pt)."""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = Pt(pt)
    run = p.add_run("")
    run.font.size = Pt(max(pt - 1, 1))


def caption(doc, text: str, align=WD_ALIGN_PARAGRAPH.CENTER):
    """Підпис під рядком: (найменування юридичної особи)."""
    return para(doc, text, size=8, align=align, italic=True)


def set_cell(cell, text: str, size: float = 10, bold: bool = False, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    if align is not None:
        p.alignment = align


def fix_widths(t, widths_cm: list[float]) -> None:
    """Фіксована ширина таблиці: Word інакше розтягує колонки за межі сторінки."""
    t.autofit = False
    tbl_pr = t._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(sum(widths_cm) * 567)))
    for i, col in enumerate(t.columns):
        col.width = Cm(widths_cm[i])
    for row in t.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)


def cell_margins(t, side_twips: int = 45) -> None:
    """Вузькі внутрішні відступи клітинок (≈0,08 см), як у PDF, — щоб слова не рвалися."""
    tbl_pr = t._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for edge, value in (("left", side_twips), ("right", side_twips), ("top", 0), ("bottom", 0)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def table(doc, rows: int, cols: int, widths_cm: list[float]):
    t = doc.add_table(rows=rows, cols=cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    fix_widths(t, widths_cm)
    cell_margins(t)
    return t


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def no_borders(t) -> None:
    tbl_pr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


def signature(doc, role: str, who: str, with_position: bool = True,
              widths: tuple[float, float, float, float] = (4.6, 5.4, 3.0, 5.0)) -> None:
    """Рядок підпису: роль — посада — ____ (підпис) — Власне ім'я ПРІЗВИЩЕ (як у pdf_out)."""
    t = doc.add_table(rows=2, cols=4)
    no_borders(t)
    fix_widths(t, list(widths))
    set_cell(t.cell(0, 0), role, size=10)
    set_cell(t.cell(0, 1), "{{ %s.position_rank }}" % who if with_position else "", size=10)
    set_cell(t.cell(0, 2), "____________", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(t.cell(0, 3), "{{ %s.name }}" % who, size=10)
    set_cell(t.cell(1, 1), "(посада, звання)" if with_position else "", size=7)
    set_cell(t.cell(1, 2), "(підпис)", size=7, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(t.cell(1, 3), "(власне ім'я та прізвище)", size=7)
    for row in t.rows:   # без зайвих відступів у клітинках
        for cell in row.cells:
            for par in cell.paragraphs:
                par.paragraph_format.space_after = Pt(0)


def processing_mark(doc, with_copies: bool = True) -> None:
    """Внизу документа: кому примірники (по рядках) і відмітка про обробку."""
    if with_copies:
        para(doc, "{{ copies_text }}", size=9)
        gap(doc, 4)
    para(doc, "{{ processing_mark }}", size=8.5, italic=True)


# --- Накладна (вимога), дод. 25 -------------------------------------------------------

def nakladna() -> Document:
    doc = new_doc()
    para(doc, "Додаток 25 до Інструкції з обліку військового майна у Збройних Силах України "
              "(пункт 24 розділу IV)", size=8, align=WD_ALIGN_PARAGRAPH.RIGHT)
    head = doc.add_table(rows=1, cols=2)
    no_borders(head)
    left, right = head.cell(0, 0), head.cell(0, 1)
    fix_widths(head, [10, 8])
    set_cell(left, "{{ unit_name }}", size=12, bold=True)
    para(left, "(найменування юридичної особи)", size=8, italic=True)
    para(left, "{% if subunit_name %}{{ subunit_name }}{% endif %}", size=11)
    para(left, "Код згідно з ЄДРПОУ {{ edrpou }}", size=11)
    set_cell(right, "Дійсна до {{ valid_until_long }}", size=11, align=WD_ALIGN_PARAGRAPH.RIGHT)

    gap(doc, 6)
    para(doc, "Накладна (вимога) № {{ doc_no }}", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "{{ place }}", align=WD_ALIGN_PARAGRAPH.CENTER)
    caption(doc, "(місце складання)")
    para(doc, "{{ doc_date_long }}", align=WD_ALIGN_PARAGRAPH.CENTER)
    caption(doc, "(дата складання)")
    para(doc, "Дата операції {{ op_date_long }}", space_after=4)
    for label, field in (("Служба забезпечення", "service"), ("Вид операції", "operation"),
                         ("Підстава (мета)", "basis"), ("Відповідальний одержувач", "recipient"),
                         ("Передає", "from_title"), ("Приймає", "to_title")):
        p = para(doc, f"{label} ")
        p.add_run("{{ %s }}" % field).underline = True
    gap(doc, 6)

    widths = [0.7, 3.8, 1.7, 1.3, 1.4, 1.7, 1.6, 1.6, 1.8, 2.4]  # разом 18 см — ширина тексту на А4 (так само в pdf_out)
    t = table(doc, 6, 10, widths)
    heads = ["№ з/п", "Назва військового майна або однорідна група (вид)", "Код номен­клатури",
             "Одиниця виміру", "Категорія (сорт)", "Вартість за одиницю", "Кількість", None, "Сума", "Примітка"]
    for i, h in enumerate(heads):
        if h:
            set_cell(t.cell(0, i), h, size=7, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    t.cell(0, 6).merge(t.cell(0, 7))
    set_cell(t.cell(0, 6), "Кількість", size=7, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(t.cell(1, 6), "відправлено (вимагається)", size=7, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(t.cell(1, 7), "прийнято (відпущено)", size=7, align=WD_ALIGN_PARAGRAPH.CENTER)
    for i in (0, 1, 2, 3, 4, 5, 8, 9):
        t.cell(0, i).merge(t.cell(1, i))
    for i in range(10):
        set_cell(t.cell(2, i), str(i + 1), size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    for r in (0, 1, 2):
        repeat_header(t.rows[r])
    set_cell(t.cell(3, 0), "{%tr for l in lines %}", size=8)
    fields = ["l.n", "l.name", "l.code", "l.uom", "l.cat", "l.price", "l.req", "l.qty", "l.sum", "l.note"]
    for i, f in enumerate(fields):
        align = WD_ALIGN_PARAGRAPH.RIGHT if f in ("l.price", "l.req", "l.qty", "l.sum") else None
        set_cell(t.cell(4, i), "{{ %s }}" % f, size=9, align=align)
    set_cell(t.cell(5, 0), "{%tr endfor %}", size=8)
    total = t.add_row()
    fix_widths(t, widths)
    total.cells[0].merge(total.cells[7])
    set_cell(total.cells[0], "Всього", size=10, bold=True)
    set_cell(total.cells[8], "{{ total_sum }}", size=10, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)

    gap(doc, 6)
    signature(doc, "Керівник (посадова особа, начальник служби)", "head")
    gap(doc, 6)
    p = para(doc, "Всього передано ")
    p.add_run("{{ total_qty_words }}").underline = True
    p.add_run(" {{ total_units_word }},")
    caption(doc, "(кількість прописом)")
    p = para(doc, "на суму ")
    p.add_run("{{ total_hrn_words }}").underline = True
    p.add_run(" грн {{ total_kop }} коп.")
    caption(doc, "(сума прописом)")
    gap(doc, 6)
    para(doc, "Матеріально відповідальні особи:")
    signature(doc, "здав:", "mvo_from")
    signature(doc, "прийняв:", "mvo_to")

    para(doc, "Відмітка фінансово-економічного органу про відображення у регістрах бухгалтерського обліку:",
         size=10)
    f = table(doc, 3, 4, [5.0, 5.0, 5.0, 3.5])
    for i, h in enumerate(["Назва облікового регістру", "За дебетом рахунку (субрахунку, коду аналітичного обліку)",
                           "За кредитом рахунку (субрахунку, коду аналітичного обліку)", "Сума"]):
        set_cell(f.cell(0, i), h, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "Особа, яка відобразила господарську операцію в бухгалтерському обліку", size=10)
    para(doc, "______________________   ____________   __________________________", size=10)
    caption(doc, "(посада)                                    (підпис)                  (власне ім'я та прізвище)",
            align=WD_ALIGN_PARAGRAPH.LEFT)
    para(doc, "«___» __________________ 20___ року", size=10)
    signature(doc, "Помічник командира з фінансово-економічної роботи — начальник фінансової служби",
              "finance", with_position=False, widths=(7.0, 3.0, 3.0, 5.0))
    processing_mark(doc)
    return doc


def line_field(doc, label: str, label_cm: float, total_cm: float = 26.7, value: str = "") -> None:
    """«Висновок ________» — підпис і лінія до правого краю (не переноситься), як у pdf_out."""
    t = doc.add_table(rows=1, cols=2)
    no_borders(t)
    fix_widths(t, [label_cm, total_cm - label_cm])
    cell_margins(t, 0)
    set_cell(t.cell(0, 0), label, size=11)
    set_cell(t.cell(0, 1), value, size=11)
    tc_pr = t.cell(0, 1)._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:color", "000000")):
        bottom.set(qn(k), v)
    borders.append(bottom)
    tc_pr.append(borders)


def _finance(doc, wide: bool) -> None:
    para(doc, "Відмітка фінансово-економічного органу про відображення у регістрах бухгалтерського обліку:",
         size=9)
    f = table(doc, 3, 4, [6.5, 7.5, 7.5, 5.2] if wide else [5.0, 5.0, 5.0, 3.0])
    for i, h in enumerate(["Назва облікового регістру", "За дебетом рахунку (субрахунку, коду аналітичного обліку)",
                           "За кредитом рахунку (субрахунку, коду аналітичного обліку)", "Сума"]):
        set_cell(f.cell(0, i), h, size=7, align=WD_ALIGN_PARAGRAPH.CENTER)
    gap(doc, 4)
    para(doc, "Особа, яка відобразила господарську операцію в бухгалтерському обліку  ______________  "
              "______________________________  «___» ____________ 20___ року", size=9)
    caption(doc, "(підпис)                              (посада, власне ім'я та прізвище)")
    gap(doc, 4)
    signature(doc, "Помічник командира з фінансово-економічної роботи — начальник фінансової служби",
              "finance", with_position=False, widths=(9.0, 3.0, 3.0, 6.0) if wide else (7.0, 3.0, 3.0, 5.0))


# --- Акти приймання-передачі, дод. 23 (ОЗ) і 24 (запаси) ---------------------------------

def act(kind: str) -> Document:
    """Альбомний А4, ширина тексту 26,7 см. Колонки й порядок блоків — як у pdf_out.act."""
    oz = kind == "oz"
    appendix, title, what = (("23", "Акт приймання-передачі основних засобів", "основні засоби") if oz else
                             ("24", "Акт приймання-передачі запасів", "запаси"))
    doc = new_doc(landscape=True)
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(1.0)
    para(doc, f"Додаток {appendix} до Інструкції з обліку військового майна у Збройних Силах України "
              "(пункт 24 розділу IV)", size=8, align=WD_ALIGN_PARAGRAPH.RIGHT)
    head = doc.add_table(rows=1, cols=2)
    no_borders(head)
    fix_widths(head, [16.7, 10])
    left, right = head.cell(0, 0), head.cell(0, 1)
    set_cell(left, "{{ unit_name }}", size=12, bold=True)
    para(left, "(найменування юридичної особи)", size=7, italic=True)
    para(left, "{% if subunit_name %}{{ subunit_name }}{% endif %}", size=11)
    para(left, "Код згідно з ЄДРПОУ {{ edrpou }}", size=11)
    set_cell(right, "ЗАТВЕРДЖУЮ", size=11, bold=True)
    para(right, "{{ approver.position }}", size=11)
    para(right, "(посада)", size=7, italic=True)
    para(right, "{{ approver.rank }} {{ approver.name }}", size=11)
    para(right, "(військове звання, власне ім'я та прізвище)", size=7, italic=True)
    para(right, "М.П. ________________________", size=11)
    para(right, "(підпис)", size=7, italic=True)
    para(right, "«___» ________________ 20___ року", size=11)
    para(doc, title, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    gap(doc, 3)

    req_w = [2.3, 2.6, 2.0, 2.8, 2.8, 5.6, 5.8, 2.8]
    req = table(doc, 2, 8, req_w)
    for i, h in enumerate(["Дата реєстрації документа", "Місце складання", "Номер документа",
                           "Дата початку приймання-передачі", "Дата закінчення приймання-передачі",
                           f"Найменування юридичної особи (ПІБ фізичної особи), що передає {what}",
                           f"Найменування юридичної особи (ПІБ фізичної особи), що приймає {what}",
                           "Служба забезпечення"]):
        set_cell(req.cell(0, i), h, size=7, align=WD_ALIGN_PARAGRAPH.CENTER)
    for i, f in enumerate(["reg_date", "place", "doc_no", "op_date", "op_date", "from_party", "to_party",
                           "service"]):
        set_cell(req.cell(1, i), "{{ %s }}" % f, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    gap(doc, 3)

    if oz:
        widths = [0.7, 4.6, 2.4, 1.3, 2.0, 1.3, 1.3, 1.3, 1.3, 2.2, 1.4, 1.5, 1.4, 2.2, 1.8]
        heads = ["№ з/п", "Назва військового майна", "Код номен­клатури", "Номер партії",
                 "Первісна (переоцінена) вартість", "Відправлено (вимагається)", None,
                 "Прийнято (відпущено)", None, "Сума", "Знос", None, "Рік випуску (побудови)",
                 "Заводський номер", "Номер паспорта"]
        subs = {5: "Кількість", 6: "Категорія", 7: "Кількість", 8: "Категорія", 10: "за одиницю", 11: "всього"}
        groups, singles, sum_col = [(5, 6), (7, 8), (10, 11)], [0, 1, 2, 3, 4, 9, 12, 13, 14], 9
        fields = ["l.n", "l.name", "l.code", "l.batch", "l.price", "l.req", "l.cat", "l.qty", "l.cat", "l.sum",
                  "l.wear_unit", "l.wear_total", "l.year_made",
                  "{% if l.serial_no %}{{ l.serial_no }}{% elif l.inventory_no %}інв. № {{ l.inventory_no }}{% endif %}",
                  "l.passport_no"]
    else:
        widths = [0.7, 6.3, 2.4, 1.6, 1.4, 2.2, 1.6, 1.3, 1.6, 1.3, 2.4, 3.9]
        heads = ["№ з/п", "Назва військового майна або однорідна група (вид)", "Код номен­клатури",
                 "Номер партії", "Одиниця виміру", "Вартість за одиницю виміру", "Відправлено (вимагається)", None,
                 "Прийнято (відпущено)", None, "Сума", "Примітки"]
        subs = {6: "Кількість", 7: "Категорія", 8: "Кількість", 9: "Категорія"}
        groups, singles, sum_col = [(6, 7), (8, 9)], [0, 1, 2, 3, 4, 5, 10, 11], 10
        fields = ["l.n", "l.name", "l.code", "l.batch", "l.uom", "l.price", "l.req", "l.cat", "l.qty", "l.cat",
                  "l.sum", "l.note"]
    n = len(widths)
    t = table(doc, 6, n, widths)
    for i, h in enumerate(heads):
        if h:
            set_cell(t.cell(0, i), h, size=7, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for i, s in subs.items():
        set_cell(t.cell(1, i), s, size=7, align=WD_ALIGN_PARAGRAPH.CENTER)
    for a, b in groups:
        title_text = heads[a]
        t.cell(0, a).merge(t.cell(0, b))
        set_cell(t.cell(0, a), title_text, size=7, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for c in singles:
        t.cell(0, c).merge(t.cell(1, c))
    for i in range(n):
        set_cell(t.cell(2, i), str(i + 1), size=7, align=WD_ALIGN_PARAGRAPH.CENTER)
    for r in (0, 1, 2):
        repeat_header(t.rows[r])
    set_cell(t.cell(3, 0), "{%tr for l in lines %}", size=7)
    for i, f in enumerate(fields):
        text = f if f.startswith("{%") else "{{ %s }}" % f
        right_align = f in ("l.price", "l.req", "l.qty", "l.sum")
        set_cell(t.cell(4, i), text, size=9, align=WD_ALIGN_PARAGRAPH.RIGHT if right_align else None)
    set_cell(t.cell(5, 0), "{%tr endfor %}", size=7)
    total = t.add_row()
    fix_widths(t, widths)
    total.cells[0].merge(total.cells[sum_col - 1])
    set_cell(total.cells[0], "Всього", size=9, bold=True)
    set_cell(total.cells[sum_col], "{{ total_sum }}", size=9, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    gap(doc, 3)

    p = para(doc, "На підставі ")
    p.add_run("{{ basis }}").underline = True
    if oz:
        p.add_run(" проведено огляд переліченого військового майна.")
        caption(doc, "(назва, дата та номер документа, на підставі якого здійснюється передача (приймання))",
                align=WD_ALIGN_PARAGRAPH.LEFT)
    else:
        p.add_run(" проведено огляд військового майна у кількості ")
        p.add_run("{{ total_qty_words }}").underline = True
        p.add_run(" {{ total_units_word }} на суму ")
        p.add_run("{{ total_hrn_words }}").underline = True
        p.add_run(" грн {{ total_kop }} коп.")
        caption(doc, "(назва, дата та номер документа, на підставі якого здійснюється передача (приймання)); "
                     "кількість і сума — прописом", align=WD_ALIGN_PARAGRAPH.LEFT)
    line_field(doc, "Місцезнаходження військового майна у момент передачі (прийняття)", 12.0)
    line_field(doc, "Військове майно визначеним вимогам відповідає / не відповідає (підкреслити необхідне)", 15.5)
    line_field(doc, "Висновок", 2.0)
    line_field(doc, "Перелік документації, що додається", 6.6)
    gap(doc, 3)
    para(doc, "Начальник служби забезпечення (начальник обліково-операційного підрозділу) військової частини "
              "(центру забезпечення):", size=9)
    signature(doc, "", "head", widths=(0.4, 9.0, 4.0, 7.0))
    para(doc, "Матеріально відповідальна особа, яка {{ our_action }} на відповідальне зберігання військове майно, "
              "вказане у цьому акті за: № ________", size=9)
    signature(doc, "", "mvo_our", widths=(0.4, 9.0, 4.0, 7.0))
    para(doc, "Військове майно прийняв / здав (підкреслити необхідне):", size=9)
    signature(doc, "", "mvo_other", widths=(0.4, 9.0, 4.0, 7.0))
    gap(doc, 3)
    _finance(doc, wide=True)
    gap(doc, 4)
    processing_mark(doc)
    return doc


# --- Наряд на видавання (приймання), дод. 5 ----------------------------------------------

def naryad() -> Document:
    """Як у частинах: книжковий А4, усе на одному аркуші, підпис командира військової частини,
    «Усього найменувань», розсилка примірників. Колонки — як у pdf_out.naryad."""
    doc = new_doc()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(1.2)
    para(doc, "Додаток 5 до Інструкції з обліку військового майна у Збройних Силах України "
              "(пункт 6 розділу III)", size=8, align=WD_ALIGN_PARAGRAPH.RIGHT)
    para(doc, "Дійсний до {{ valid_until_long }}")
    para(doc, "НАРЯД № {{ doc_no }}", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "на видавання (приймання) військового майна", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    gap(doc, 4)

    def req(heads, fields, widths):
        t = table(doc, 2, len(heads), widths)
        for i, h in enumerate(heads):
            set_cell(t.cell(0, i), h, size=7, align=WD_ALIGN_PARAGRAPH.CENTER)
        for i, f in enumerate(fields):
            set_cell(t.cell(1, i), f, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        gap(doc, 4)

    req(["Реєстраційний номер", "Номер аркуша", "Номер документа", "Дата документа", "Підстава (мета) операції",
         "Дата операції", "Служба забезпечення органу військового управління"],
        ["{{ reg_no }}", "", "{{ doc_no }}", "{{ doc_date }}", "{{ basis }}", "{{ op_date }}", "{{ service }}"],
        [1.7, 1.3, 1.6, 1.8, 6.0, 1.8, 3.8])
    req(["Вантажовідправник", "Вантажоодержувач та його поштова адреса", "Вид транспорту", "Номер транспорту",
         "Найменування транспортного документа", "Номер транспортного документа"],
        ["{{ from_title }}", "{{ to_title }}{% if to_address %}, {{ to_address }}{% endif %}",
         "{{ order.transport_kind }}", "{{ order.transport_no }}", "{{ order.transport_doc_name }}",
         "{{ order.transport_doc_no }}"],
        [3.6, 5.4, 2.0, 2.0, 2.8, 2.2])

    widths = [0.7, 3.9, 1.8, 1.2, 1.7, 1.3, 1.75, 1.75, 2.4, 1.5]   # 18 см
    heads = ["№ з/п", "Найменування військового майна (індекс, номер креслення)", "Код номенклатури",
             "Одиниця виміру", "Ціна за одиницю", "Категорія (сорт)",
             "Кількість до видавання (приймання), відвантаження",
             "Фактично видано (прийнято), відвантажено",
             "Назва, номер та дата документа, за яким здійснено видавання (приймання), відвантаження", "Примітка"]
    t = table(doc, 5, 10, widths)
    for i, h in enumerate(heads):
        # Звичайний (не жирний) шрифт: вузькі графи книжкового аркуша, інакше Word рве слова.
        set_cell(t.cell(0, i), h, size=7, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell(t.cell(1, i), str(i + 1), size=7, align=WD_ALIGN_PARAGRAPH.CENTER)
    repeat_header(t.rows[0])
    repeat_header(t.rows[1])
    set_cell(t.cell(2, 0), "{%tr for l in lines %}", size=7)
    fields = ["l.n", "l.name", "l.code", "l.uom", "l.price", "l.cat", "l.qty", "l.fact", "l.fact_docs", "l.note"]
    for i, f in enumerate(fields):
        right = f in ("l.price", "l.qty", "l.fact")
        set_cell(t.cell(3, i), "{{ %s }}" % f, size=9, align=WD_ALIGN_PARAGRAPH.RIGHT if right else None)
    set_cell(t.cell(4, 0), "{%tr endfor %}", size=7)
    total = t.add_row()
    fix_widths(t, widths)
    set_cell(total.cells[1], "Усього найменувань: {{ count }} ({{ count_words }})", size=9, bold=True)
    gap(doc, 8)

    # Те, що за формою на звороті, — на тому ж аркуші (як у частинах).
    line_field(doc, "Порядок відправлення", 4.0, total_cm=18.0, value="{{ order.dispatch_order }}")
    gap(doc, 4)
    g = doc.add_table(rows=1, cols=4)
    no_borders(g)
    fix_widths(g, [6.5, 3.5, 3.8, 4.2])
    set_cell(g.cell(0, 0), "Строк прибуття варти з приймальником", size=11)
    set_cell(g.cell(0, 1), "{{ order.guard_term }}", size=11)
    set_cell(g.cell(0, 2), "від військової частини", size=11)
    set_cell(g.cell(0, 3), "{{ order.guard_unit }}", size=11)
    gap(doc, 14)

    # Підпис командира військової частини: посада і звання ліворуч, «Ім'я ПРІЗВИЩЕ» праворуч.
    s = doc.add_table(rows=2, cols=3)
    no_borders(s)
    fix_widths(s, [8.0, 4.0, 6.0])
    set_cell(s.cell(0, 0), "{{ approver.position }}", size=11)
    set_cell(s.cell(0, 1), "____________", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(s.cell(0, 2), "{{ approver.name }}", size=11, align=WD_ALIGN_PARAGRAPH.RIGHT)
    set_cell(s.cell(1, 0), "{{ approver.rank }}", size=11)
    set_cell(s.cell(1, 1), "(підпис)", size=7, align=WD_ALIGN_PARAGRAPH.CENTER)
    gap(doc, 12)

    # Унизу: М.П., розсилка примірників — ліворуч; таблиця упаковки — праворуч.
    b = doc.add_table(rows=1, cols=2)
    no_borders(b)
    fix_widths(b, [10.5, 7.5])
    left = b.cell(0, 0)
    set_cell(left, "М. П.  «___» ________________ 20___ року", size=11)
    para(left, "")
    para(left, "{{ copies_text }}", size=11)
    right = b.cell(0, 1)
    right.text = ""
    pack = right.add_table(rows=6, cols=2)
    pack.style = "Table Grid"
    fix_widths(pack, [5.0, 2.3])
    for i, h in enumerate(("Упаковано місць", "Вид (характер) упаковки", "Маркування", "Маса",
                           "Відправлено місць", "Дата відправлення")):
        set_cell(pack.cell(i, 0), h, size=11)
    gap(doc, 8)
    processing_mark(doc, with_copies=False)   # примірники — ліворуч унизу
    return doc


# --- Акт якісного (технічного) стану, дод. 1 до Порядку списання (наказ МОУ № 81) -------------

def condition() -> Document:
    """Альбомний А4, 17 граф — ширини як у pdf_out.CONDITION_WIDTHS."""
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
    from oblik.pdf_out import CONDITION_HEADS, CONDITION_REQ, CONDITION_REQ_WIDTHS, CONDITION_WIDTHS

    doc = new_doc(landscape=True)
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(1.0)
    para(doc, "Додаток 1 до Порядку списання військового майна у Збройних Силах України "
              "(пункт 1.5 розділу I)", size=8, align=WD_ALIGN_PARAGRAPH.RIGHT)
    head = doc.add_table(rows=1, cols=2)
    no_borders(head)
    fix_widths(head, [16.7, 10])
    left, right = head.cell(0, 0), head.cell(0, 1)
    acc = left.add_table(rows=2, cols=3)
    acc.style = "Table Grid"
    fix_widths(acc, [3.6, 3.0, 3.6])
    for i, h in enumerate(["Номенклатурний номер", "Основний рахунок", "Кореспондентський рахунок"]):
        set_cell(acc.cell(0, i), h, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(right, "ЗАТВЕРДЖУЮ", size=11, bold=True)
    para(right, "{{ approver.position }}", size=11)
    para(right, "{{ approver.rank }} {{ approver.name }}", size=11)
    para(right, "(посада, військове звання, власне ім'я та прізвище)", size=7, italic=True)
    para(right, "«___» ________________ 20___ року", size=11)
    para(right, "М.П.", size=11)
    para(doc, "Акт якісного (технічного) стану № {{ doc_no }}", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "{{ cond.group_name }}", size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    gap(doc, 3)

    req = table(doc, 3, len(CONDITION_REQ), CONDITION_REQ_WIDTHS)
    for i, (title, code, field) in enumerate(CONDITION_REQ):
        set_cell(req.cell(0, i), title, size=7, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell(req.cell(1, i), code, size=7, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell(req.cell(2, i), "{{ %s }}" % field if field else "", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    gap(doc, 3)
    para(doc, "У результаті огляду встановлено:", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)

    heads = CONDITION_HEADS
    t = table(doc, 7, 17, CONDITION_WIDTHS)
    center = WD_ALIGN_PARAGRAPH.CENTER
    for i, h in enumerate(heads):
        if h and i and i != 8:
            set_cell(t.cell(1, i), h, size=7, align=center)
    set_cell(t.cell(2, 8), "за нормою", size=7, align=center)
    set_cell(t.cell(2, 9), "фак­тично", size=7, align=center)
    for i in range(17):
        set_cell(t.cell(3, i), str(i + 1), size=7, align=center)
    # Об'єднання клітинок шапки (як у pdf_out.condition).
    t.cell(0, 1).merge(t.cell(0, 9))
    set_cell(t.cell(0, 1), "Списати", size=8, bold=True, align=center)
    t.cell(0, 10).merge(t.cell(0, 16))
    set_cell(t.cell(0, 10), "Оприбуткувати", size=8, bold=True, align=center)
    t.cell(1, 8).merge(t.cell(1, 9))
    set_cell(t.cell(1, 8), "експлуатується, років", size=7, align=center)
    for col in list(range(1, 8)) + list(range(10, 17)):
        text = heads[col]
        t.cell(1, col).merge(t.cell(2, col))
        set_cell(t.cell(1, col), text, size=7, align=center)
    t.cell(0, 0).merge(t.cell(2, 0))
    set_cell(t.cell(0, 0), "№ з/п", size=7, align=center)
    for r in range(4):
        repeat_header(t.rows[r])
    set_cell(t.cell(4, 0), "{%tr for r in cond.rows %}", size=7)
    fields = ["r.n", "r.name", "r.code", "r.uom", "r.cat", "r.qty", "r.price", "r.sum", "r.ynorm", "r.yfact",
              "r.in_name", "r.in_code", "r.in_uom", "r.in_cat", "r.in_qty", "r.in_price", "r.in_sum"]
    for i, f in enumerate(fields):
        right_align = f.split(".")[1] in ("qty", "price", "sum", "in_qty", "in_price", "in_sum")
        set_cell(t.cell(5, i), "{{ %s }}" % f, size=9,
                 align=WD_ALIGN_PARAGRAPH.RIGHT if right_align else (None if "name" in f else center))
    set_cell(t.cell(6, 0), "{%tr endfor %}", size=7)
    total = t.add_row()
    fix_widths(t, CONDITION_WIDTHS)
    set_cell(total.cells[1], "Усього", size=9, bold=True, align=center)
    for col, field in ((5, "cond.out_qty"), (7, "cond.out_sum"), (14, "cond.in_qty"), (16, "cond.in_sum")):
        set_cell(total.cells[col], "{{ %s }}" % field, size=9, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    gap(doc, 6)

    para(doc, "Висновок комісії: {{ cond.commission_conclusion }}", size=11)
    gap(doc, 6)
    para(doc, "Голова комісії:", size=11)
    signature(doc, "", "cond.head", widths=(0.4, 11.0, 4.0, 7.0))
    gap(doc, 3)
    para(doc, "Члени комісії:", size=11)
    para(doc, "{%p for m in cond.members %}", size=2)
    signature(doc, "", "m", widths=(0.4, 11.0, 4.0, 7.0))
    para(doc, "{%p endfor %}", size=2)
    gap(doc, 4)
    para(doc, "Висновок старшого начальника: {{ cond.senior_conclusion }}", size=11)
    signature(doc, "", "cond.senior", widths=(0.4, 11.0, 4.0, 7.0))
    para(doc, "«___» ________________ 20___ року          М.П.", size=11)
    gap(doc, 6)
    para(doc, "Отримані від розбирання вузли, прилади, запасні частини, деталі та інше майно, вказане в графах "
              "11–17, на відповідальне зберігання прийняв", size=11)
    signature(doc, "", "cond.keeper", widths=(0.4, 11.0, 4.0, 7.0))
    para(doc, "«___» ________________ 20___ року", size=11)
    gap(doc, 4)
    processing_mark(doc)
    return doc


TEMPLATES = {"nakladna.docx": nakladna,
             "act_condition.docx": condition,
             "naryad.docx": naryad,
             "act_oz.docx": lambda: act("oz"),
             "act_zap.docx": lambda: act("zap")}


def main() -> int:
    """Без аргументів — усі шаблони; інакше лише названі (nakladna.docx act_oz.docx …).
    Генеруйте лише той, що змінили: файл .docx містить час створення, тож навіть без змін
    у вмісті він стає «іншим», і програма вважатиме його новою стандартною версією."""
    OUT.mkdir(parents=True, exist_ok=True)
    wanted = sys.argv[1:] or list(TEMPLATES)
    for name, build in TEMPLATES.items():
        if name not in wanted:
            continue
        build().save(OUT / name)
        print("створено", OUT / name)
    return 0


if __name__ == "__main__":
    sys.exit(main())

"""Звіти: залишки на дату (по МВО, підрозділу, особі), картка руху одиниці.

Усе рахується з журналу руху (balances.py); тут — лише відбір, підписи й оформлення.
Функції нічого не змінюють у базі.
"""
from __future__ import annotations

import sqlite3
from dataclasses import dataclass, field
from io import BytesIO

from . import balances, documents, locations, refs, settings
from .documents import CATEGORY_TITLES, DOC_TYPES
from .textutil import QTY_SCALE, format_money, format_qty, uk_sort_key

CLASS_TITLES = {"fixed": "Основні засоби", "low_value": "МНМА", "inventory": "Запаси",
                "mshp": "МШП", "off_balance": "Забалансове"}


class ReportError(ValueError):
    pass


# --- Залишки на дату -------------------------------------------------------------

@dataclass
class StockReport:
    title: str                       # «Склад ТЗО (разом з виданим особам)»
    on_date: str
    rows: list[dict] = field(default_factory=list)
    total_kop: int = 0
    many_places: bool = False        # показувати колонку «Де»


def scope_choices(conn: sqlite3.Connection) -> list[tuple[str, list[tuple[str, str]]]]:
    """Варіанти «чиї залишки»: група → [(код, підпис)]."""
    labels = refs.Labels(conn)
    units = [(f"u:{r['id']}", labels.get("units", r["id"])) for r in conn.execute(
        "SELECT id FROM units WHERE is_accounting = 1 ORDER BY name COLLATE UK")]
    mvo = [(f"m:{r['mvo_person_id']}", labels.get("persons", r["mvo_person_id"])) for r in conn.execute(
        "SELECT DISTINCT mvo_person_id FROM units WHERE is_accounting = 1 AND mvo_person_id IS NOT NULL")]
    mvo.sort(key=lambda o: uk_sort_key(o[1]))
    persons = [(f"p:{pid}", title) for pid, title in refs.options(conn, "persons")]
    # Підрозділи, що самі не є місцем обліку, але містять місця обліку (наприклад, відділення
    # з двома МВО або підпорядкована частина) — разом усе, що в них.
    groups = []
    for row in conn.execute("SELECT id FROM units WHERE is_accounting = 0 AND is_active = 1 "
                            "ORDER BY name COLLATE UK"):
        inner = locations.subtree_units(conn, row["id"])
        n = conn.execute(f"SELECT COUNT(*) FROM units WHERE is_accounting = 1 AND id IN "
                         f"({','.join('?' * len(inner))})", inner).fetchone()[0]
        if n:
            groups.append((f"t:{row['id']}", f"{labels.get('units', row['id'])} — разом {n} місць обліку"
                           if n > 1 else f"{labels.get('units', row['id'])} — разом"))
    return [("Уся частина", [("all", "Усі місця обліку й особи")]),
            ("МВО (усі його місця обліку)", mvo),
            ("Підрозділ разом з усіма місцями обліку в ньому", groups),
            ("Місце обліку (склад, підрозділ)", units),
            ("Особа (видане в користування)", persons)]


def resolve_scope(conn, scope: str, include_subunits: bool) -> tuple[list[int] | None, str]:
    """Код вибору → (місця для balances, заголовок)."""
    labels = refs.Labels(conn)
    if scope in ("", "all"):
        return None, "Уся частина"
    kind, _, raw = scope.partition(":")
    if not raw.isdigit():
        raise ReportError("Оберіть, чиї залишки показати")
    ref_id = int(raw)
    if kind == "u":
        if conn.execute("SELECT 1 FROM units WHERE id = ? AND is_accounting = 1", (ref_id,)).fetchone() is None:
            raise ReportError("Місце обліку не знайдено")
        title = labels.get("units", ref_id) + " (разом з виданим особам"
        title += ", з підлеглими підрозділами)" if include_subunits else ")"
        return locations.scope_for_unit(conn, ref_id, include_subunits), title
    if kind == "t":
        if conn.execute("SELECT 1 FROM units WHERE id = ?", (ref_id,)).fetchone() is None:
            raise ReportError("Підрозділ не знайдено")
        return (locations.scope_for_unit(conn, ref_id, include_subunits=True),
                labels.get("units", ref_id) + " (усі місця обліку в ньому)")
    if kind == "m":
        unit_ids = [r["id"] for r in conn.execute(
            "SELECT id FROM units WHERE is_accounting = 1 AND mvo_person_id = ?", (ref_id,))]
        if not unit_ids:
            raise ReportError("Ця особа не є МВО жодного місця обліку")
        locs: list[int] = []
        for u in unit_ids:
            locs += [x for x in locations.scope_for_unit(conn, u, include_subunits) if x not in locs]
        return locs, "МВО " + labels.get("persons", ref_id)
    if kind == "p":
        if conn.execute("SELECT 1 FROM persons WHERE id = ?", (ref_id,)).fetchone() is None:
            raise ReportError("Особу не знайдено")
        return locations.person_locations(conn, ref_id), "Видане особі: " + labels.get("persons", ref_id)
    raise ReportError("Оберіть, чиї залишки показати")


def stock(conn: sqlite3.Connection, scope: str, on_date: str,
          include_subunits: bool = False) -> StockReport:
    location_ids, title = resolve_scope(conn, scope, include_subunits)
    noms = {r["id"]: r for r in conn.execute("SELECT * FROM nomenclature")}
    items = {r["id"]: r for r in conn.execute("SELECT * FROM items")}
    place_titles: dict[int, str] = {}
    report = StockReport(title=title, on_date=on_date)
    for b in balances.balances(conn, location_ids, on_date):
        nom = noms[b["nomenclature_id"]]
        item = items.get(b["item_id"]) if b["item_id"] else None
        if b["location_id"] not in place_titles:
            place_titles[b["location_id"]] = locations.title(conn, b["location_id"])
        report.rows.append({
            **b,
            "name": nom["name"],
            "code": nom["code"] or "",
            "uom": nom["uom"],
            "class_title": CLASS_TITLES.get(nom["accounting_class"], ""),
            "serial_no": (item["serial_no"] or "") if item else "",
            "inventory_no": (item["inventory_no"] or "") if item else "",
            "place": place_titles[b["location_id"]],
            "cat": CATEGORY_TITLES.get(b["category"], ""),
        })
        report.total_kop += b["value_kop"]
    report.rows.sort(key=lambda r: (uk_sort_key(r["name"]), r["inventory_no"], r["serial_no"],
                                    r["category"] or 0, r["price_kop"], uk_sort_key(r["place"])))
    report.many_places = location_ids is None or len(location_ids) > 1
    for n, r in enumerate(report.rows, 1):
        r["n"] = n
    return report


def stock_xlsx(conn: sqlite3.Connection, report: StockReport) -> bytes:
    """Залишки у форматі Excel (відкривається і в LibreOffice Calc)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter

    s = settings.load(conn)
    wb = Workbook()
    ws = wb.active
    ws.title = "Залишки"
    heads = ["№ з/п", "Найменування", "Код номенклатури", "Заводський №", "Інвентарний №",
             "Од. виміру", "Категорія", "Ціна, грн", "Кількість", "Сума, грн"]
    widths = [6, 42, 16, 16, 16, 9, 10, 13, 11, 15]
    if report.many_places:
        heads.insert(1, "Де перебуває")
        widths.insert(1, 28)

    ws.append([s.get("unit_name") or ""])
    if s.get("subunit_name"):
        ws.append([s["subunit_name"]])
    ws.append([f"Залишки майна на {documents.fmt_date(report.on_date)}"])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=13)
    ws.append([report.title])
    ws.append([])
    ws.append(heads)
    head_row = ws.max_row
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in ws[head_row]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = border

    for r in report.rows:
        values = [r["n"], r["name"], r["code"], r["serial_no"], r["inventory_no"], r["uom"], r["cat"],
                  r["price_kop"] / 100, r["qty_m"] / QTY_SCALE, r["value_kop"] / 100]
        if report.many_places:
            values.insert(1, r["place"])
        ws.append(values)
        for c in ws[ws.max_row]:
            c.border = border
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ncols = len(heads)
    money_cols = (ncols - 2, ncols)   # ціна, сума
    qty_col = ncols - 1
    for row in ws.iter_rows(min_row=head_row + 1, max_row=ws.max_row):
        for c in money_cols:
            row[c - 1].number_format = "# ##0.00"
        row[qty_col - 1].number_format = "# ##0.###"

    ws.append([])
    total = ws.max_row + 1
    ws.cell(total, 2, f"Усього найменувань (рядків): {len(report.rows)}").font = Font(bold=True)
    ws.cell(total, ncols - 1, "Разом:").font = Font(bold=True)
    sum_cell = ws.cell(total, ncols, report.total_kop / 100)
    sum_cell.font = Font(bold=True)
    sum_cell.number_format = "# ##0.00"

    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(head_row + 1, 1)
    ws.print_title_rows = f"{head_row}:{head_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- Картка руху поштучної одиниці ------------------------------------------------

def item_card(conn: sqlite3.Connection, item_id: int) -> dict:
    item = conn.execute("SELECT * FROM items WHERE id = ?", (item_id,)).fetchone()
    if item is None:
        raise ReportError("Одиницю не знайдено")
    labels = refs.Labels(conn)
    nom = conn.execute("SELECT * FROM nomenclature WHERE id = ?", (item["nomenclature_id"],)).fetchone()
    moves = []
    for m in balances.item_card(conn, item_id):
        moves.append({
            **m,
            "date": documents.fmt_date(m["op_date"]),
            "doc_label": documents.label(conn, m["document_id"]),
            "type_title": DOC_TYPES[m["doc_type"]].title,
            "from_title": locations.title(conn, m["from_location_id"]),
            "to_title": locations.title(conn, m["to_location_id"]),
            "cat": CATEGORY_TITLES.get(m["category"], ""),
            "price": format_money(m["price_kop"]),
            "qty": format_qty(m["qty_m"]),
        })
    pos = balances.item_position(conn, item_id)
    components = conn.execute("SELECT * FROM item_components WHERE item_id = ? ORDER BY line_no",
                              (item_id,)).fetchall()
    return {
        "item": item,
        "label": labels.get("items", item_id),
        "nom": nom,
        "class_title": CLASS_TITLES.get(nom["accounting_class"], ""),
        "moves": moves,
        "position": {
            "place": locations.title(conn, pos["location_id"]),
            "cat": CATEGORY_TITLES.get(pos["category"], ""),
            "price": format_money(pos["price_kop"]),
        } if pos else None,
        "components": [{**dict(c), "qty": format_qty(c["qty_m"])} for c in components],
        "initial_cost": format_money(item["initial_cost_kop"]),
    }


def find_items(conn: sqlite3.Connection, q: str, limit: int = 200) -> list[dict]:
    """Пошук одиниць за номером або найменуванням, з поточним місцем."""
    labels = refs.Labels(conn)
    where, params = "", []
    if q:
        like = f"%{q.lower()}%"
        where = ("WHERE ulower(i.serial_no) LIKE ? OR ulower(i.inventory_no) LIKE ? "
                 "OR ulower(i.passport_no) LIKE ? OR ulower(n.name) LIKE ?")
        params = [like] * 4
    rows = conn.execute(
        "SELECT i.id, i.serial_no, i.inventory_no, n.name FROM items i "
        f"JOIN nomenclature n ON n.id = i.nomenclature_id {where} "
        "ORDER BY n.name COLLATE UK, i.inventory_no, i.serial_no LIMIT ?", [*params, limit]).fetchall()
    result = []
    for r in rows:
        pos = balances.item_position(conn, r["id"])
        result.append({**dict(r), "label": labels.get("items", r["id"]),
                       "place": locations.title(conn, pos["location_id"]) if pos else "не на обліку"})
    return result
